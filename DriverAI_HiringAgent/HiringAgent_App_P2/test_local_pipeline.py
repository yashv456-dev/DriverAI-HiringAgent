"""Offline reliability regressions. All databases, CVs and publications are synthetic."""
import hashlib
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from hiring_agent.store import SQLiteCandidateStore, RowVanished
from hiring_agent.sqlite_store import database_path
from hiring_agent.local_pipeline import run_local_pipeline, LocalClient, _child, worker_lock
from hiring_agent.publisher import publish_results
from sharepoint_client import SharePointError, SharePointClient


def candidate(app='APP-A', date='2026-09-08T10:00:00'):
    return {'Application ID': app, 'Email': app.lower()+'@example.test',
            'Full Name': 'Alex Morgan', 'Received Date': date, 'Last Updated Date': date,
            'Application Updates': 0, 'Original Filename': 'resume.docx', 'Has Resume': 'Yes',
            'Status': 'New Email Received', 'Mail Body': '', 'Mail Subject': 'Application',
            'Resume URL': 'manifest:'+json.dumps([{'name': app+'.docx', 'folder': '/CVs'}])}


class Remote:
    _wb_path = '/Master_Files/Sharepoint_Master_File.xlsx'
    resumes_folder = '/CVs'
    def __init__(self, rows=None):
        self.rows = rows or [candidate()]
        self.rejected = []
        self.files = {('/CVs', r['Application ID']+'.docx'): b'original '+r['Application ID'].encode() for r in self.rows}
        self.uploads = []
        self.fail_upload = False
    def list_rows(self):
        return [{'index': i, 'values': dict(r)} for i, r in enumerate(self.rows)]
    def list_rejected_rows(self):
        return [{'index': i, 'values': dict(r)} for i, r in enumerate(self.rejected)]
    def download_file(self, folder, name):
        key = ('/'+folder.strip('/'), name)
        if key not in self.files:
            raise SharePointError('missing', status_code=404)
        return self.files[key]
    def file_web_url(self, folder, name):
        return 'https://example.test'+folder+'/'+name
    def upload_file(self, folder, name, data):
        if self.fail_upload:
            raise SharePointError('locked', status_code=423)
        self.uploads.append((folder, name))
        self.files[('/'+folder.strip('/'), name)] = data


class ReliabilityTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        env = patch.dict(os.environ, {'HIRING_SQLITE_PATH': str(self.root/'candidates.db'),
                                    'HIRING_PUBLISH_DIR': str(self.root/'results'),
                                    'HIRING_MASTER_REPORT_PATH': '/Master_Files/P2-MasterFile.xlsx',
                                    'HIRING_CLIENT_REPORT_PATH': '/Candidate_List_Results.xlsx'})
        env.start(); self.addCleanup(env.stop)
        roles = patch('hiring_agent.jd_sources.get_active_roles', return_value=[{'title': 'Engineer', 'skills': ['Python']}])
        roles.start(); self.addCleanup(roles.stop)
        delay = patch('hiring_agent.config.ROW_DELAY_SECONDS', 0)
        delay.start(); self.addCleanup(delay.stop)
        self.store = SQLiteCandidateStore(':memory:')
        self.addCleanup(self.store.conn.close)

    @staticmethod
    def score(values, docs, roles):
        values = dict(values, Status='Scored', Country='United States', Category='Engineering')
        return {'summary': {'processed': 1}, 'rows': [('main', values)]}

    def test_incremental_import_keeps_completed_and_imports_new(self):
        remote = Remote()
        self.assertEqual(self.store.sync_from_client(remote), 1)
        self.store.save_by_id('APP-A', {'Status': 'Scored'})
        remote.rows.append(candidate('APP-B'))
        self.assertEqual(self.store.sync_from_client(remote), 1)
        self.assertEqual(self.store.get('APP-A').values['Status'], 'Scored')
        self.assertIsNotNone(self.store.get('APP-B'))
        self.assertEqual(self.store.sync_from_client(remote, force=True), 0)

    def test_new_input_requeues_and_preserves_result_history(self):
        remote = Remote()
        self.store.sync_from_client(remote)
        self.store.save_by_id('APP-A', {'Status': 'Scored', 'Phone': 'old'})
        remote.rows[0]['Last Updated Date'] = '2026-09-09T10:00:00'
        remote.rows[0]['Application Updates'] = 1
        self.store.sync_from_client(remote)
        row = self.store.get('APP-A').values
        self.assertEqual(row['Status'], 'New Email Received')
        self.assertNotIn('Phone', row)
        self.assertGreater(self.store.conn.execute('SELECT count(*) FROM candidate_history').fetchone()[0], 0)

    def test_one_application_is_one_live_row_on_one_sheet(self):
        """A candidate is on CandidateList or Rejected - never both, and never twice.

        This used to be defended only at write time (a hinted resolution picking between two
        live rows). It is now a UNIQUE index, so the doubled state cannot be created at all -
        which is the property that matters, because the old Excel backend produced it whenever
        Graph refused the delete half of a rejection and left the candidate on both sheets.
        """
        import sqlite3
        self.store.add(dict(candidate(), label='first'))
        with self.assertRaises(sqlite3.IntegrityError):
            self.store.add(dict(candidate(), label='duplicate'))
        self.assertEqual(len(self.store.all_rows()), 1)

        # Rejection MOVES the single row; it must not leave a copy behind on the main sheet.
        self.store.move_to_rejected('APP-A', {'Status': 'Rejected - Non-USA Location'})
        self.assertEqual([r.app_id for r in self.store.all_rows()], [])
        self.assertEqual([r.app_id for r in self.store.all_rows('rejected')], ['APP-A'])
        live = self.store.conn.execute(
            "SELECT COUNT(*) FROM candidates WHERE active=1 AND app_id='APP-A'").fetchone()[0]
        self.assertEqual(live, 1)

        # A hint that names no live row still fails closed rather than guessing a neighbour.
        with self.assertRaises(RowVanished):
            self.store.save_by_id('APP-A', {}, hint=999)

    def test_conflicting_identity_is_retained_not_merged(self):
        remote = Remote([candidate(), dict(candidate(), Email='someoneelse@example.test')])
        self.store.sync_from_client(remote)
        self.assertEqual(len(self.store.all_rows()), 0)
        self.assertEqual(self.store.conn.execute('SELECT count(*) FROM source_events').fetchone()[0], 2)
        self.assertEqual(self.store.conn.execute('SELECT count(*) FROM import_conflicts').fetchone()[0], 1)

    def test_stale_worker_cannot_commit_newer_input(self):
        self.store.add(candidate())
        row = self.store.get('APP-A')
        attempt, version = self.store.begin_attempt(int(row.version))
        self.store.save_by_id('APP-A', {'Mail Body': 'new input'})
        state = self.store.finish_attempt(attempt, version, ('main', dict(candidate(), Status='Scored')))
        self.assertEqual(state, 'superseded')
        self.assertEqual(self.store.get('APP-A').values['Mail Body'], 'new input')

    def test_pipeline_rerun_does_not_rescore_or_touch_intake(self):
        remote = Remote()
        with patch('hiring_agent.local_pipeline.isolated_score'):
            first = run_local_pipeline(remote, scorer=self.score)
        self.assertEqual(first['processed'], 1)
        self.assertEqual(remote.rows[0]['Status'], 'New Email Received')
        second = run_local_pipeline(remote, scorer=lambda *a: self.fail('completed input scored twice'))
        self.assertEqual(second['processed'], 0)
        self.assertFalse(any(name == 'Sharepoint_Master_File.xlsx' for _, name in remote.uploads))
        wb = load_workbook(self.root/'results'/'generation_2'/'P2-MasterFile.xlsx')
        self.assertEqual(wb.sheetnames, ['CandidateList', 'Rejected'])
        self.assertTrue(all(ws.tables for ws in wb)); wb.close()

    def test_upload_failure_retries_without_scoring(self):
        remote = Remote(); remote.fail_upload = True
        result = run_local_pipeline(remote, scorer=self.score)
        self.assertEqual(result['processed'], 1)
        store = SQLiteCandidateStore(database_path())
        self.assertNotEqual(*store.conn.execute('SELECT revision,published_revision FROM publication').fetchone())
        store.conn.close()
        remote.fail_upload = False
        result = run_local_pipeline(remote, scorer=lambda *a: self.fail('publication retried inference'))
        self.assertEqual(result['processed'], 0)
        store = SQLiteCandidateStore(database_path())
        revision, published = store.conn.execute(
            'SELECT revision,published_revision FROM publication').fetchone()
        store.conn.close()
        self.assertEqual(revision, published)
        # P2 owns exactly two remote files and overwrites them in place. Versioned '.r<N>'
        # siblings and a remote pointer file used to be published alongside, which is how the
        # tenant accumulated four 'master' workbooks a human had to choose between; version
        # history belongs in the local generation_<N> folders, and published_revision above
        # already answers 'what is live'.
        self.assertEqual(set(remote.uploads),
                         {('Master_Files', 'P2-MasterFile.xlsx'),
                          ('', 'Candidate_List_Results.xlsx')})

    def test_bad_candidate_does_not_block_next_or_consume_retry(self):
        remote = Remote([candidate(), candidate('APP-B')])
        def score(values, docs, roles):
            if values['Application ID'] == 'APP-A': raise TimeoutError('model slow')
            return self.score(values, docs, roles)
        result = run_local_pipeline(remote, scorer=score)
        self.assertEqual((result['processed'], result['deferred']), (1, 1))
        store = SQLiteCandidateStore(database_path())
        self.assertEqual(store.get('APP-A').values.get('Retry Count', 0), 0)
        store.conn.close()

    def test_missing_slot_defers_without_partial_score(self):
        remote = Remote()
        remote.rows[0]['Resume URL'] = 'manifest:'+json.dumps([{'name': 'APP-A.docx', 'folder': '/CVs'}, {'name': 'missing.docx', 'folder': '/CVs'}])
        result = run_local_pipeline(remote, scorer=lambda *a: self.fail('partial attachments scored'))
        self.assertEqual(result['deferred'], 1)

    def test_exporter_refuses_intake_target(self):
        self.store.add(candidate())
        remote = Remote()
        with patch.dict(os.environ, {'HIRING_MASTER_REPORT_PATH': remote._wb_path}):
            with self.assertRaises(ValueError): publish_results(self.store, remote)
        self.assertEqual(remote.uploads, [])

    def test_copy_preserves_original_and_rejects_collision(self):
        remote = Remote()
        remote.delete_file = lambda *a: self.fail('original CV deleted')
        self.assertTrue(SharePointClient.move_resume(remote, 'APP-A.docx', '/CVs', '/Rejected'))
        self.assertIn(('/CVs', 'APP-A.docx'), remote.files)
        remote.files[('/Rejected', 'APP-A.docx')] = b'different CV'
        with self.assertRaises(SharePointError): SharePointClient.move_resume(remote, 'APP-A.docx', '/CVs', '/Rejected')

    def test_initializer_preserves_existing_database(self):
        import init_db
        with patch.object(init_db, 'DB_PATH', self.root/'init.db'):
            store = init_db.init_empty_db(); store.add(candidate()); store.conn.close()
            store = init_db.init_empty_db()
            self.assertIsNotNone(store.get('APP-A')); store.conn.close()

    def test_backup_restore(self):
        self.store.add(candidate())
        self.store.backup(self.root/'backup.db')
        restored = SQLiteCandidateStore(self.root/'backup.db')
        self.assertEqual(restored.get('APP-A').values, candidate()); restored.conn.close()

    def test_existing_scoring_logic_runs_only_against_local_adapter(self):
        from hiring_agent import sharepoint_scoring as scoring
        values = candidate()
        path = self.root/'cv.docx'; path.write_bytes(b'synthetic')
        docs = {'APP-A.docx': {'path': str(path), 'sha256': hashlib.sha256(b'synthetic').hexdigest(),
                              'url': 'https://example.test/cv.docx'}}
        values['_Local Resume Names'] = list(docs)
        values['Resume URL'] = docs['APP-A.docx']['url']
        details = {'full_name': 'Alex Morgan', 'phone': '5125550123', 'location': 'Austin, TX',
                   'country': 'United States', 'skills': 'Python', 'looking_for_role': 'Engineer',
                   'education': 'BS Computer Science', 'experience': '3 years'}
        capture = []
        class Pipe:
            def send(self, value): capture.append(value)
            def close(self): pass
        with patch.object(scoring, 'ollama_health', return_value=(True,'OK')), \
             patch.object(scoring, 'extract_text_from_bytes', return_value='Alex Morgan Austin TX United States Python'), \
             patch.object(scoring, 'extract_candidate_details_smart', return_value=details), \
             patch.object(scoring, 'ai_recheck_fields', side_effect=lambda fields,*a: fields), \
             patch.object(scoring, 'infer_missing_portfolios', return_value=('N/A','N/A','N/A')), \
             patch.object(scoring, 'suggested_roles', return_value={'role_1':'Engineer (90%)','role_2':'','role_3':'','source':'ollama'}), \
             patch.dict(scoring.EXTRACTION_SOURCE, {'value':'ollama'}), \
             patch.dict(os.environ, {'HIRING_P2_DISABLED':'false'}):
            _child(Pipe(), values, docs, [{'title':'Engineer','skills':['Python']}])
        self.assertNotIn('error', capture[0], capture[0])
        self.assertEqual(capture[0]['summary']['errors'], 0, capture[0])
        self.assertEqual(capture[0]['rows'][0][1]['Status'], 'Scored', capture[0])


if __name__ == '__main__':
    unittest.main()
