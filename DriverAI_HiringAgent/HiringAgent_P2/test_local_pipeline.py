"""Offline reliability regressions. All databases, CVs and publications are synthetic."""
import hashlib
import io
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

import sys
sys.path.insert(0, str(Path(__file__).resolve().parent))
import test_isolation  # noqa: F401 - must precede every hiring_agent import

from hiring_agent.store import SQLiteCandidateStore, RowVanished
from hiring_agent.sqlite_store import database_path
from hiring_agent.local_pipeline import (run_local_pipeline, LocalClient, _child, worker_lock,
                                         _find_by_app_ref, cache_documents)
from hiring_agent.publisher import publish_results, _verify_report
from sharepoint_client import SharePointError, SharePointClient


def candidate(app='APP-A', date='2026-09-08T10:00:00'):
    return {'Application ID': app, 'Email': app.lower()+'@example.test',
            'Full Name': 'Alex Morgan', 'Received Date': date, 'Last Updated Date': date,
            'Application Updates': 0, 'Original Filename': 'resume.docx', 'Has Resume': 'Yes',
            'Status': 'New Email Received', 'Mail Body': '', 'Mail Subject': 'Application',
            'Resume URL': 'manifest:'+json.dumps([{'name': app+'.docx', 'folder': '/CVs'}])}


class Remote:
    _wb_path = '/Master_Files/Sharepoint_Master_File.xlsx'
    resumes_folder = '/CVs'
    def __init__(self, rows=None):
        self.rows = rows or [candidate()]
        self.rejected = []
        self.files = {('/CVs', r['Application ID']+'.docx'): b'original '+r['Application ID'].encode() for r in self.rows}
        self.uploads = []
        self.fail_upload = False
        self.fail_upload_name = None
    def list_rows(self):
        return [{'index': i, 'values': dict(r)} for i, r in enumerate(self.rows)]
    def list_rejected_rows(self):
        return [{'index': i, 'values': dict(r)} for i, r in enumerate(self.rejected)]
    def download_file(self, folder, name):
        key = ('/'+folder.strip('/'), name)
        if key not in self.files:
            raise SharePointError('missing', status_code=404)
        return self.files[key]
    def file_web_url(self, folder, name):
        return 'https://example.test'+folder+'/'+name
    def upload_file(self, folder, name, data):
        if self.fail_upload or name == self.fail_upload_name:
            raise SharePointError('locked', status_code=423)
        self.uploads.append((folder, name))
        self.files[('/'+folder.strip('/'), name)] = data
    def delete_file(self, folder, name):
        self.files.pop(('/'+folder.strip('/'), name), None)


class ReliabilityTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        env = patch.dict(os.environ, {'HIRING_SQLITE_PATH': str(self.root/'candidates.db'),
                                    'HIRING_PUBLISH_DIR': str(self.root/'results'),
                                    'HIRING_MASTER_REPORT_PATH': '/Master_Files/P2-MasterFile.xlsx',
                                    'HIRING_CLIENT_REPORT_PATH': '/Candidate_List_Results.xlsx'})
        env.start(); self.addCleanup(env.stop)
        roles = patch('hiring_agent.jd_sources.get_active_roles', return_value=[{'title': 'Engineer', 'skills': ['Python']}])
        roles.start(); self.addCleanup(roles.stop)
        delay = patch('hiring_agent.config.ROW_DELAY_SECONDS', 0)
        delay.start(); self.addCleanup(delay.stop)
        self.store = SQLiteCandidateStore(':memory:')
        self.addCleanup(self.store.conn.close)

    @staticmethod
    def score(values, docs, roles):
        values = dict(values, Status='Scored', Country='United States', Category='Engineering')
        return {'summary': {'processed': 1}, 'rows': [('main', values)]}

    def test_incremental_import_keeps_completed_and_imports_new(self):
        remote = Remote()
        self.assertEqual(self.store.sync_from_client(remote), 1)
        self.store.save_by_id('APP-A', {'Status': 'Scored'})
        remote.rows.append(candidate('APP-B'))
        self.assertEqual(self.store.sync_from_client(remote), 1)
        self.assertEqual(self.store.get('APP-A').values['Status'], 'Scored')
        self.assertIsNotNone(self.store.get('APP-B'))
        self.assertEqual(self.store.sync_from_client(remote, force=True), 0)

    def test_new_input_requeues_and_preserves_result_history(self):
        remote = Remote()
        self.store.sync_from_client(remote)
        self.store.save_by_id('APP-A', {'Status': 'Scored', 'Phone': 'old'})
        remote.rows[0]['Last Updated Date'] = '2026-09-09T10:00:00'
        remote.rows[0]['Application Updates'] = 1
        self.store.sync_from_client(remote)
        row = self.store.get('APP-A').values
        self.assertEqual(row['Status'], 'New Email Received')
        self.assertNotIn('Phone', row)
        self.assertGreater(self.store.conn.execute('SELECT count(*) FROM candidate_history').fetchone()[0], 0)

    def test_one_application_is_one_live_row_on_one_sheet(self):
        """A candidate is on CandidateList or Rejected - never both, and never twice.

        This used to be defended only at write time (a hinted resolution picking between two
        live rows). It is now a UNIQUE index, so the doubled state cannot be created at all -
        which is the property that matters, because the old Excel backend produced it whenever
        Graph refused the delete half of a rejection and left the candidate on both sheets.
        """
        import sqlite3
        self.store.add(dict(candidate(), label='first'))
        with self.assertRaises(sqlite3.IntegrityError):
            self.store.add(dict(candidate(), label='duplicate'))
        self.assertEqual(len(self.store.all_rows()), 1)

        # Rejection MOVES the single row; it must not leave a copy behind on the main sheet.
        self.store.move_to_rejected('APP-A', {'Status': 'Rejected - Non-USA Location'})
        self.assertEqual([r.app_id for r in self.store.all_rows()], [])
        self.assertEqual([r.app_id for r in self.store.all_rows('rejected')], ['APP-A'])
        live = self.store.conn.execute(
            "SELECT COUNT(*) FROM candidates WHERE active=1 AND app_id='APP-A'").fetchone()[0]
        self.assertEqual(live, 1)

        # A hint that names no live row still fails closed rather than guessing a neighbour.
        with self.assertRaises(RowVanished):
            self.store.save_by_id('APP-A', {}, hint=999)

    def test_conflicting_identity_is_retained_not_merged(self):
        remote = Remote([candidate(), dict(candidate(), Email='someoneelse@example.test')])
        with self.assertRaisesRegex(RuntimeError, 'unresolved Application ID conflict'):
            self.store.sync_from_client(remote)
        self.assertEqual(len(self.store.all_rows()), 0)
        self.assertEqual(self.store.conn.execute('SELECT count(*) FROM source_events').fetchone()[0], 2)
        self.assertEqual(self.store.conn.execute('SELECT count(*) FROM import_conflicts').fetchone()[0], 1)

    def test_candidate_shaped_p1_row_without_id_stops_sync(self):
        broken = dict(candidate(), **{'Application ID': ''})
        with self.assertRaisesRegex(RuntimeError, 'without an Application ID'):
            self.store.sync_from_client(Remote([broken]))
        self.assertEqual(len(self.store.all_rows()), 0)

    def test_stale_worker_cannot_commit_newer_input(self):
        self.store.add(candidate())
        row = self.store.get('APP-A')
        attempt, version = self.store.begin_attempt(int(row.version))
        self.store.save_by_id('APP-A', {'Mail Body': 'new input'})
        state = self.store.finish_attempt(attempt, version, ('main', dict(candidate(), Status='Scored')))
        self.assertEqual(state, 'superseded')
        self.assertEqual(self.store.get('APP-A').values['Mail Body'], 'new input')

    def test_pipeline_rerun_does_not_rescore_or_touch_intake(self):
        remote = Remote()
        with patch('hiring_agent.local_pipeline.isolated_score'):
            first = run_local_pipeline(remote, scorer=self.score)
        self.assertEqual(first['processed'], 1)
        self.assertEqual(remote.rows[0]['Status'], 'New Email Received')
        second = run_local_pipeline(remote, scorer=lambda *a: self.fail('completed input scored twice'))
        self.assertEqual(second['processed'], 0)
        self.assertFalse(any(name == 'Sharepoint_Master_File.xlsx' for _, name in remote.uploads))
        wb = load_workbook(self.root/'results'/'generation_2'/'P2-MasterFile.xlsx')
        self.assertEqual(wb.sheetnames, ['CandidateList', 'Rejected', 'Unfamiliar Role List'])
        self.assertTrue(all(ws.tables for ws in wb)); wb.close()

    def _publish_and_open(self, remote):
        """Publish the current store and return the master workbook."""
        from hiring_agent.publisher import publish_results
        store = SQLiteCandidateStore(database_path())
        self.addCleanup(store.conn.close)
        store.sync_from_client(remote)
        out = Path(str(publish_results(store, remote, upload=False)))
        master = next(Path(out).parent.rglob('P2-MasterFile.xlsx')) if out.is_dir() else \
            next(out.parent.rglob('P2-MasterFile.xlsx'))
        return store, load_workbook(master)

    @staticmethod
    def _ids(ws):
        headers = [c.value for c in ws[1]]
        col = headers.index('Application ID') + 1
        return [str(ws.cell(i, col).value) for i in range(2, ws.max_row + 1)
                if ws.cell(i, col).value]

    def test_no_matching_role_candidate_is_listed_on_its_own_sheet(self):
        """A candidate with no opening belongs on 'Unfamiliar Role List', and only there.

        Live: Ruta Kothari and Kshitij Sahu, both silicon engineers against a catalogue with
        no chip-design role. Listing them on CandidateList beside genuinely matched people
        misrepresents them; leaving them off the workbook entirely loses them.
        """
        from hiring_agent import config as cfg
        remote = Remote([candidate('APP-MATCHED'), candidate('APP-NOMATCH')])
        store, wb = self._publish_and_open(remote)
        store.save_by_id('APP-MATCHED', {'Status': 'Scored', 'Category': 'Engineering'})
        store.save_by_id('APP-NOMATCH', {'Status': cfg.STATUS_NO_MATCHING_ROLE,
                                         'Category': 'General',
                                         'Suggested Role 1': cfg.NO_ROLE_MATCH_LABEL})
        wb.close()
        _, wb = self._publish_and_open(remote)
        try:
            self.assertEqual(wb.sheetnames,
                             ['CandidateList', 'Rejected', 'Unfamiliar Role List'])
            self.assertEqual(self._ids(wb['CandidateList']), ['APP-MATCHED'])
            self.assertEqual(self._ids(wb['Unfamiliar Role List']), ['APP-NOMATCH'])
            # Every sheet stays a real Excel table, so the workbook opens the same way.
            self.assertTrue(all(ws.tables for ws in wb))
        finally:
            wb.close()

    def test_unfamiliar_sheet_is_published_even_when_empty(self):
        """The schema is fixed, so a run with nobody in it still writes the sheet."""
        remote = Remote()
        _, wb = self._publish_and_open(remote)
        try:
            self.assertIn('Unfamiliar Role List', wb.sheetnames)
            self.assertEqual(self._ids(wb['Unfamiliar Role List']), [])
        finally:
            wb.close()

    def test_upload_failure_retries_without_scoring(self):
        remote = Remote(); remote.fail_upload = True
        result = run_local_pipeline(remote, scorer=self.score)
        self.assertEqual(result['processed'], 1)
        store = SQLiteCandidateStore(database_path())
        self.assertNotEqual(*store.conn.execute('SELECT revision,published_revision FROM publication').fetchone())
        store.conn.close()
        remote.fail_upload = False
        result = run_local_pipeline(remote, scorer=lambda *a: self.fail('publication retried inference'))
        self.assertEqual(result['processed'], 0)
        store = SQLiteCandidateStore(database_path())
        revision, published = store.conn.execute(
            'SELECT revision,published_revision FROM publication').fetchone()
        store.conn.close()
        self.assertEqual(revision, published)
        # P2 owns exactly two remote files and overwrites them in place. Versioned '.r<N>'
        # siblings and a remote pointer file used to be published alongside, which is how the
        # tenant accumulated four 'master' workbooks a human had to choose between; version
        # history belongs in the local generation_<N> folders, and published_revision above
        # already answers 'what is live'.
        self.assertEqual(set(remote.uploads),
                         {('Master_Files', 'P2-MasterFile.xlsx'),
                          ('', 'Candidate_List_Results.xlsx')})

    def test_second_workbook_failure_rolls_back_first_upload(self):
        row = dict(candidate(), Status='Scored', Country='United States',
                   Category='Engineering', Phone='(602) 555-0100')
        self.store.add(row)
        remote = Remote()
        publish_results(self.store, remote)
        old_client = remote.files[('/', 'Candidate_List_Results.xlsx')]
        old_master = remote.files[('/Master_Files', 'P2-MasterFile.xlsx')]
        old_published = self.store.conn.execute(
            'SELECT published_revision FROM publication WHERE id=1').fetchone()[0]

        self.store.save_by_id('APP-A', {'Phone': '(602) 555-0199'})
        remote.fail_upload_name = 'P2-MasterFile.xlsx'
        publish_results(self.store, remote)

        self.assertEqual(remote.files[('/', 'Candidate_List_Results.xlsx')], old_client)
        self.assertEqual(remote.files[('/Master_Files', 'P2-MasterFile.xlsx')], old_master)
        publication = self.store.conn.execute(
            'SELECT published_revision,error FROM publication WHERE id=1').fetchone()
        self.assertEqual(publication['published_revision'], old_published)
        self.assertIn('rolled back successfully', publication['error'])

    def test_remote_verifier_rejects_same_shape_wrong_candidate(self):
        row = dict(candidate(), Status='Scored', Country='United States',
                   Category='Engineering', Phone='(602) 555-0100')
        self.store.add(row)
        remote = Remote()
        report = publish_results(self.store, remote)
        expected = report.read_bytes()
        wb = load_workbook(io.BytesIO(expected))
        ws = wb['Candidates']
        id_col = [cell.value for cell in ws[1]].index('Application ID') + 1
        ws.cell(3, id_col, 'APP-WRONG')
        changed = io.BytesIO()
        wb.save(changed)
        wb.close()
        remote.files[('/', 'Candidate_List_Results.xlsx')] = changed.getvalue()
        with self.assertRaisesRegex(ValueError, 'missing IDs=.*APP-A'):
            _verify_report(remote, '', 'Candidate_List_Results.xlsx', expected)

    def test_result_audit_rejects_omitted_and_changed_rows(self):
        from hiring_agent.sharepoint_scoring import (
            audit_client_export_integrity, prepare_client_export_rows)
        source = [dict(candidate('APP-A'), Category='Engineering'),
                  dict(candidate('APP-B'), Category='Engineering')]
        prepared = prepare_client_export_rows(source)
        omitted = [dict(row) for row in prepared[:-1]]
        with self.assertRaisesRegex(ValueError, 'omitted source Application IDs'):
            audit_client_export_integrity(omitted, source)

        changed = [dict(row) for row in prepared]
        changed[1]['Email'] = 'wrong@example.test'
        with self.assertRaisesRegex(ValueError, 'differs from committed source'):
            audit_client_export_integrity(changed, source)

    def test_bad_candidate_does_not_block_next_or_consume_retry(self):
        remote = Remote([candidate(), candidate('APP-B')])
        def score(values, docs, roles):
            if values['Application ID'] == 'APP-A': raise TimeoutError('model slow')
            return self.score(values, docs, roles)
        result = run_local_pipeline(remote, scorer=score)
        self.assertEqual((result['processed'], result['deferred']), (1, 1))
        store = SQLiteCandidateStore(database_path())
        self.assertEqual(store.get('APP-A').values.get('Retry Count', 0), 0)
        store.conn.close()

    def test_row_stops_being_retried_once_its_budget_is_spent(self):
        """A deterministic failure must not be re-attempted forever.

        Before the cap reached this path, APP-20260821-0812-MCQA accumulated 22 deferred
        attempts against a SCORE_RETRY_MAX of 3: the give-up logic lived in the excel
        scorer, which the sqlite pipeline never reaches.
        """
        remote = Remote()
        def always_fails(values, docs, roles):
            raise TimeoutError('model slow')

        with patch('hiring_agent.config.SCORE_RETRY_MAX', 3):
            for run in range(1, 4):
                result = run_local_pipeline(remote, scorer=always_fails)
                self.assertEqual(result['deferred'], 1, f'run {run} should have attempted')
            # Budget spent: the row is skipped instead of burning another attempt.
            result = run_local_pipeline(remote, scorer=always_fails)
            self.assertEqual(result['deferred'], 0)

        store = SQLiteCandidateStore(database_path())
        self.addCleanup(store.conn.close)
        row = store.get('APP-A')
        self.assertEqual(store.deferred_attempts(int(row.version)), 3,
                         'the capped run must not record a fourth attempt')
        # Never silently rejected or altered - it stays put for a human to look at.
        self.assertEqual(row.values['Status'], 'New Email Received')

    def test_new_input_restores_the_retry_budget(self):
        """A candidate who sends a fresh resume is not punished for the old one."""
        remote = Remote()
        def always_fails(values, docs, roles):
            raise TimeoutError('model slow')

        with patch('hiring_agent.config.SCORE_RETRY_MAX', 1):
            self.assertEqual(run_local_pipeline(remote, scorer=always_fails)['deferred'], 1)
            self.assertEqual(run_local_pipeline(remote, scorer=always_fails)['deferred'], 0)
            remote.rows[0]['Last Updated Date'] = '2026-09-09T10:00:00'
            remote.rows[0]['Application Updates'] = 1
            self.assertEqual(run_local_pipeline(remote, scorer=always_fails)['deferred'], 1,
                             'new input must earn a fresh attempt')

    def test_explicit_rerun_overrides_a_spent_budget(self):
        """An operator asking for one specific row by id is not blocked by the cap."""
        remote = Remote()
        def always_fails(values, docs, roles):
            raise TimeoutError('model slow')

        with patch('hiring_agent.config.SCORE_RETRY_MAX', 1):
            self.assertEqual(run_local_pipeline(remote, scorer=always_fails)['deferred'], 1)
            self.assertEqual(run_local_pipeline(remote, scorer=always_fails)['deferred'], 0)
            forced = run_local_pipeline(remote, scorer=always_fails, app_ids=['APP-A'])
            self.assertEqual(forced['deferred'], 1)

    def test_missing_slot_defers_without_partial_score(self):
        remote = Remote()
        remote.rows[0]['Resume URL'] = 'manifest:'+json.dumps([{'name': 'APP-A.docx', 'folder': '/CVs'}, {'name': 'missing.docx', 'folder': '/CVs'}])
        result = run_local_pipeline(remote, scorer=lambda *a: self.fail('partial attachments scored'))
        self.assertEqual(result['deferred'], 1)

    def test_short_tail_collision_uses_original_filename_identity(self):
        """Two unrelated Application IDs may share the same four-character tail."""
        class TailRemote:
            def list_folder_children(self, _place):
                return [{'name': 'Shalin_Edward_PMCA.pdf'},
                        {'name': 'Vibha_Swaminathan_Resume_-_DriverAI_PMCA.pdf'}]
            def download_file(self, _place, name):
                return name.encode()

        found = _find_by_app_ref(
            TailRemote(), 'APP-20260805-0931-PMCA', ['/CVs'],
            hints=('Vibha_Swaminathan Resume - DriverAI.pdf', 'vibha.work.1212'))
        self.assertEqual(found[1], 'Vibha_Swaminathan_Resume_-_DriverAI_PMCA.pdf')

    def test_short_tail_collision_fails_closed_without_identity_match(self):
        class TailRemote:
            def list_folder_children(self, _place):
                return [{'name': 'Shalin_Edward_PMCA.pdf'},
                        {'name': 'Vibha_Swaminathan_PMCA.pdf'}]
            def download_file(self, _place, name):
                return name.encode()

        self.assertIsNone(_find_by_app_ref(
            TailRemote(), 'APP-20260805-0931-PMCA', ['/CVs'], hints=('someone_else.pdf',)))

    def test_cache_prefers_original_filename_over_stale_derived_link(self):
        """A prior bad result must not keep selecting the other PMCA candidate forever."""
        class TailRemote:
            resumes_folder = '/Candidate_Resumes'
            files = {
                ('/Candidate_Resumes/2026/August', 'Shalin_Edward_PMCA.pdf'): b'shalin',
                ('/Candidate_Resumes/2026/August',
                 'Vibha_Swaminathan_Resume_-_DriverAI_PMCA.pdf'): b'vibha',
            }
            def download_file(self, place, name):
                key = ('/' + place.strip('/'), name)
                if key not in self.files:
                    raise SharePointError('missing', status_code=404)
                return self.files[key]
            def file_web_url(self, place, name):
                return 'https://example.test/' + name
            def list_folder_children(self, place):
                prefix = '/' + place.strip('/')
                return [{'name': name} for folder, name in self.files if folder == prefix]

        values = {
            'Application ID': 'APP-20260805-0931-PMCA',
            'Received Date': '2026-08-05T09:31:49',
            'Original Filename': 'Vibha_Swaminathan Resume - DriverAI.pdf',
            'Full Name': 'Shalin Edward',
            'Email': 'vibha.work.1212@gmail.com',
            'Category': 'Data Analytics',
            'Resume URL': 'https://example.test/Shalin_Edward_PMCA.pdf',
        }
        docs = cache_documents(TailRemote(), values, self.root/'tail-cache')
        self.assertEqual(list(docs), ['Vibha_Swaminathan_Resume_-_DriverAI_PMCA.pdf'])
        self.assertEqual(Path(next(iter(docs.values()))['path']).read_bytes(), b'vibha')

    def test_exporter_refuses_intake_target(self):
        self.store.add(candidate())
        remote = Remote()
        with patch.dict(os.environ, {'HIRING_MASTER_REPORT_PATH': remote._wb_path}):
            with self.assertRaises(ValueError): publish_results(self.store, remote)
        self.assertEqual(remote.uploads, [])

    def test_copy_preserves_original_and_rejects_collision(self):
        remote = Remote()
        remote.delete_file = lambda *a: self.fail('original CV deleted')
        self.assertTrue(SharePointClient.move_resume(remote, 'APP-A.docx', '/CVs', '/Rejected'))
        self.assertIn(('/CVs', 'APP-A.docx'), remote.files)
        remote.files[('/Rejected', 'APP-A.docx')] = b'different CV'
        with self.assertRaises(SharePointError): SharePointClient.move_resume(remote, 'APP-A.docx', '/CVs', '/Rejected')

    def test_initializer_preserves_existing_database(self):
        import init_db
        with patch.object(init_db, 'DB_PATH', self.root/'init.db'):
            store = init_db.init_empty_db(); store.add(candidate()); store.conn.close()
            store = init_db.init_empty_db()
            self.assertIsNotNone(store.get('APP-A')); store.conn.close()

    def test_backup_restore(self):
        self.store.add(candidate())
        self.store.backup(self.root/'backup.db')
        restored = SQLiteCandidateStore(self.root/'backup.db')
        self.assertEqual(restored.get('APP-A').values, candidate()); restored.conn.close()

    def test_existing_scoring_logic_runs_only_against_local_adapter(self):
        from hiring_agent import sharepoint_scoring as scoring
        values = candidate()
        path = self.root/'cv.docx'; path.write_bytes(b'synthetic')
        docs = {'APP-A.docx': {'path': str(path), 'sha256': hashlib.sha256(b'synthetic').hexdigest(),
                              'url': 'https://example.test/cv.docx'}}
        values['_Local Resume Names'] = list(docs)
        values['Resume URL'] = docs['APP-A.docx']['url']
        details = {'full_name': 'Alex Morgan', 'phone': '5125550123', 'location': 'Austin, TX',
                   'country': 'United States', 'skills': 'Python', 'looking_for_role': 'Engineer',
                   'education': 'BS Computer Science', 'experience': '3 years'}
        capture = []
        class Pipe:
            def send(self, value): capture.append(value)
            def close(self): pass
        with patch.object(scoring, 'ollama_health', return_value=(True,'OK')), \
             patch.object(scoring, 'extract_text_from_bytes', return_value='Alex Morgan Austin TX United States Python'), \
             patch.object(scoring, 'extract_candidate_details_smart', return_value=details), \
             patch.object(scoring, 'ai_recheck_fields', side_effect=lambda fields,*a: fields), \
             patch.object(scoring, 'infer_missing_portfolios', return_value=('N/A','N/A','N/A')), \
             patch.object(scoring, 'suggested_roles', return_value={'role_1':'Engineer (90%)','role_2':'','role_3':'','source':'ollama'}), \
             patch.dict(scoring.EXTRACTION_SOURCE, {'value':'ollama'}), \
             patch.dict(os.environ, {'HIRING_P2_DISABLED':'false'}):
            _child(Pipe(), values, docs, [{'title':'Engineer','skills':['Python']}])
        self.assertNotIn('error', capture[0], capture[0])
        self.assertEqual(capture[0]['summary']['errors'], 0, capture[0])
        self.assertEqual(capture[0]['rows'][0][1]['Status'], 'Scored', capture[0])

    def test_candidate_with_no_matching_opening_is_recorded_not_retried(self):
        """A read CV that matches no JD is an answer, not a failure.

        APP-20260915-1110-OP7A (VLSI engineer, catalogue has no chip-design opening) was
        deferred every run and so never reached the master file at all - not scored, not
        rejected, invisible to anyone reading the sheet. She belongs on Main under General
        for a human to place, and NOT on the Rejected sheet, which queues an automated
        decline for a candidate nobody decided to turn down.
        """
        from hiring_agent import sharepoint_scoring as scoring
        from hiring_agent import config as cfg
        values = candidate()
        path = self.root/'cv.docx'; path.write_bytes(b'synthetic')
        docs = {'APP-A.docx': {'path': str(path), 'sha256': hashlib.sha256(b'synthetic').hexdigest(),
                              'url': 'https://example.test/cv.docx'}}
        values['_Local Resume Names'] = list(docs)
        values['Resume URL'] = docs['APP-A.docx']['url']
        details = {'full_name': 'Alex Morgan', 'phone': '5125550123', 'location': 'Austin, TX',
                   'country': 'United States', 'skills': 'Verilog, ASIC, RTL',
                   'looking_for_role': 'EE Circuit Design Engineer',
                   'education': 'MS Computer Engineering', 'experience': '3 years'}
        no_match = {'role_1': '', 'role_2': '', 'role_3': '', 'source': 'no_match',
                    'reason': 'Ollama ranked 9 role(s) but none cleared the scoring filters'}
        capture = []
        class Pipe:
            def send(self, value): capture.append(value)
            def close(self): pass
        with patch.object(scoring, 'ollama_health', return_value=(True,'OK')), \
             patch.object(scoring, 'extract_text_from_bytes', return_value='Alex Morgan Austin TX Verilog ASIC RTL'), \
             patch.object(scoring, 'extract_candidate_details_smart', return_value=details), \
             patch.object(scoring, 'ai_recheck_fields', side_effect=lambda fields,*a: fields), \
             patch.object(scoring, 'infer_missing_portfolios', return_value=('N/A','N/A','N/A')), \
             patch.object(scoring, 'suggested_roles', return_value=no_match), \
             patch.dict(scoring.EXTRACTION_SOURCE, {'value':'ollama'}), \
             patch.dict(os.environ, {'HIRING_P2_DISABLED':'false'}):
            _child(Pipe(), values, docs, [{'title':'Engineer','skills':['Python']}])

        self.assertNotIn('error', capture[0], capture[0])
        summary = capture[0]['summary']
        self.assertEqual(summary.get('deferred', 0), 0,
                         'a no-match candidate must not be queued for another attempt')
        sheet, row = capture[0]['rows'][0]
        self.assertEqual(sheet, 'main', 'no-match belongs on Main, never the Rejected sheet')
        self.assertEqual(row['Status'], cfg.STATUS_NO_MATCHING_ROLE)
        self.assertNotEqual(row['Status'], cfg.STATUS_NEEDS_REVIEW,
                            'must not be reported as an unreadable resume - it read fine')
        self.assertEqual(row['Category'], 'General')
        self.assertEqual(row['Suggested Role 1'], cfg.NO_ROLE_MATCH_LABEL,
                         'no role cleared the bar, so the row must say so outright')
        self.assertEqual(row['Suggested Role 2'], cfg.NO_ROLE_MATCH_LABEL)
        self.assertEqual(row['Suggested Role 3'], cfg.NO_ROLE_MATCH_LABEL)
        # Her extracted profile has to survive - that is the point of recording her.
        self.assertEqual(row['Full Name'], 'Alex Morgan')
        self.assertIn('Verilog', row['Current Skills'])

    def test_match_under_the_placement_bar_is_not_published_as_a_match(self):
        """A thin percentage is not a placement.

        Scores drift run to run near the bottom of the range - the same CV produced 0% and
        35% on consecutive runs - so a weak match is not a stable verdict, and publishing
        one files a candidate under a category they do not belong to. Below
        scoring.publish_min_percent the row reads 'Not Matching' under General, exactly like
        a candidate whose field has no opening at all.
        """
        from hiring_agent import sharepoint_scoring as scoring
        from hiring_agent import config as cfg
        values = candidate()
        path = self.root/'cv.docx'; path.write_bytes(b'synthetic')
        docs = {'APP-A.docx': {'path': str(path), 'sha256': hashlib.sha256(b'synthetic').hexdigest(),
                              'url': 'https://example.test/cv.docx'}}
        values['_Local Resume Names'] = list(docs)
        values['Resume URL'] = docs['APP-A.docx']['url']
        details = {'full_name': 'Alex Morgan', 'phone': '5125550123', 'location': 'Austin, TX',
                   'country': 'United States', 'skills': 'Python, SQL',
                   'looking_for_role': 'Data Analyst',
                   'education': 'BS Computer Science', 'experience': '3 years'}
        weak = {'role_1': 'Data Analyst (55%)', 'role_2': 'Business Data Analyst (40%)',
                'role_3': '', 'source': 'ollama'}
        capture = []
        class Pipe:
            def send(self, value): capture.append(value)
            def close(self): pass
        with patch.object(scoring, 'ollama_health', return_value=(True,'OK')), \
             patch.object(scoring, 'extract_text_from_bytes', return_value='Alex Morgan Austin TX Python SQL'), \
             patch.object(scoring, 'extract_candidate_details_smart', return_value=details), \
             patch.object(scoring, 'ai_recheck_fields', side_effect=lambda fields,*a: fields), \
             patch.object(scoring, 'infer_missing_portfolios', return_value=('N/A','N/A','N/A')), \
             patch.object(scoring, 'suggested_roles', return_value=weak), \
             patch.dict(scoring.EXTRACTION_SOURCE, {'value':'ollama'}), \
             patch.dict(os.environ, {'HIRING_P2_DISABLED':'false'}):
            _child(Pipe(), values, docs, [{'title':'Engineer','skills':['Python']}])

        self.assertNotIn('error', capture[0], capture[0])
        sheet, row = capture[0]['rows'][0]
        self.assertEqual(sheet, 'main')
        self.assertEqual(row['Status'], cfg.STATUS_NO_MATCHING_ROLE)
        self.assertEqual(row['Suggested Role 1'], cfg.NO_ROLE_MATCH_LABEL)
        self.assertEqual(row['Category'], 'General',
                         'Category derives from Role 1, so it must not keep Data Analytics')

    def test_non_usa_is_rejected_even_when_nothing_matched(self):
        """Outside the USA goes to Rejected, match or no match.

        Geography is settled before the placement bar, so a candidate who is both non-USA
        and unmatched is a rejection - not an Unfamiliar Role List entry. That sheet is for
        people we could hire but have no opening for.
        """
        from hiring_agent import sharepoint_scoring as scoring
        values = candidate()
        path = self.root/'cv.docx'; path.write_bytes(b'synthetic')
        docs = {'APP-A.docx': {'path': str(path), 'sha256': hashlib.sha256(b'synthetic').hexdigest(),
                              'url': 'https://example.test/cv.docx'}}
        values['_Local Resume Names'] = list(docs)
        values['Resume URL'] = docs['APP-A.docx']['url']
        details = {'full_name': 'Alex Morgan', 'phone': '442012345678',
                   'location': 'London, United Kingdom', 'country': 'United Kingdom',
                   'skills': 'Python, SQL', 'looking_for_role': 'Data Analyst',
                   'education': 'BS Computer Science', 'experience': '3 years'}
        weak = {'role_1': 'Data Analyst (55%)', 'role_2': '', 'role_3': '', 'source': 'ollama'}
        capture = []
        class Pipe:
            def send(self, value): capture.append(value)
            def close(self): pass
        with patch.object(scoring, 'ollama_health', return_value=(True,'OK')), \
             patch.object(scoring, 'extract_text_from_bytes', return_value='Alex Morgan London United Kingdom'), \
             patch.object(scoring, 'extract_candidate_details_smart', return_value=details), \
             patch.object(scoring, 'ai_recheck_fields', side_effect=lambda fields,*a: fields), \
             patch.object(scoring, 'infer_missing_portfolios', return_value=('N/A','N/A','N/A')), \
             patch.object(scoring, 'suggested_roles', return_value=weak), \
             patch.dict(scoring.EXTRACTION_SOURCE, {'value':'ollama'}), \
             patch.dict(os.environ, {'HIRING_P2_DISABLED':'false'}):
            _child(Pipe(), values, docs, [{'title':'Engineer','skills':['Python']}])

        self.assertNotIn('error', capture[0], capture[0])
        sheet, row = capture[0]['rows'][0]
        self.assertEqual(sheet, 'rejected',
                         'non-USA outranks the placement bar; this is a rejection')

    def test_no_matching_role_row_is_not_picked_up_again(self):
        """Parked on a human, not on the bot: re-running must not re-score it.

        The answer cannot change until the catalogue does, so re-queueing it would burn a
        model call every run and rewrite the row for ever - the loop this whole status
        exists to end. `--recheck-row` is the way back in once a matching JD lands.
        """
        from hiring_agent import config as cfg
        remote = Remote()
        store = SQLiteCandidateStore(database_path())
        self.addCleanup(store.conn.close)
        store.sync_from_client(remote)
        store.save_by_id('APP-A', {'Status': cfg.STATUS_NO_MATCHING_ROLE,
                                   'Category': 'General', 'Suggested Role 1': ''})

        result = run_local_pipeline(
            remote, scorer=lambda *a: self.fail('a parked no-match row must not be re-scored'))
        self.assertEqual(result['processed'], 0)
        self.assertEqual(store.get('APP-A').values['Status'], cfg.STATUS_NO_MATCHING_ROLE)

        # ...but an explicit re-run still reaches it.
        forced = run_local_pipeline(remote, scorer=self.score, app_ids=['APP-A'])
        self.assertEqual(forced['processed'], 1)


class MasterWriteContract(unittest.TestCase):
    """P2 may write only MASTER_WRITE_COLUMNS onto P1's intake workbook."""

    P1_ROW = {'Application ID': 'APP-1', 'Email': 'a@b.com',
              'Received Date': '2026-09-01T10:00:00',
              'Last Updated Date': '2026-09-01T10:00:00', 'Application Updates': 0,
              'Original Filename': '', 'Mail Subject': 'Application', 'Mail Body': 'hi',
              'Has Resume': 'Yes', 'Status': 'New Email Received'}

    def _scored(self):
        from hiring_agent.extraction import MISSING_VALUE
        row = dict(self.P1_ROW)
        row.update({'Status': 'Scored', 'Phone': '(602) 555-0147',
                    'Resume Link': 'Marcus_Vance_AF6A', 'Resume URL': 'https://x/y.pdf',
                    'Resume Folder Path': '/a/b', 'Category': 'General',
                    'Education': MISSING_VALUE,
                    # the two input columns P2 fills when P1 leaves them blank
                    'Original Filename': 'resume.pdf',
                    'Last Updated Date': '2026-09-02T12:00:00'})
        return row

    def test_main_patch_carries_only_the_two_agreed_columns(self):
        from hiring_agent.local_pipeline import _master_patch, MASTER_WRITE_COLUMNS
        patch_body = _master_patch(self._scored())
        self.assertEqual(set(patch_body), set(MASTER_WRITE_COLUMNS))
        self.assertEqual(patch_body['Status'], 'Scored')
        self.assertEqual(patch_body['Resume Link'], 'Marcus_Vance_AF6A')

    def test_main_patch_never_carries_a_p1_input_column(self):
        from hiring_agent.local_pipeline import _master_patch
        from hiring_agent.sqlite_store import INPUT_COLUMNS
        body = _master_patch(self._scored())
        self.assertFalse(set(body) & set(INPUT_COLUMNS),
                         'an input column in a master write resets the row on the next poll')

    def test_resume_is_filed_in_the_folder_of_its_sheet(self):
        """A rejection puts the CV in the year's Rejected folder; a return puts it back.

        The local pipeline recorded wherever it found the file, so a candidate rejected on a
        re-score kept '/Candidate_Resumes/2026/August' on the Rejected sheet (2026-09-17).
        """
        from hiring_agent.local_pipeline import _file_by_sheet

        class Mover:
            resumes_folder = '/CVs'
            def __init__(self, fail=False):
                self.moves, self.fail = [], fail
            def move_resume(self, name, from_folder, to_folder):
                if self.fail:
                    raise SharePointError('locked', status_code=423)
                self.moves.append((name, from_folder, to_folder))
                return True
            def file_web_url(self, folder, name):
                return 'https://example.test' + folder + '/' + name

        received = '2026-08-05T00:35:50'
        in_month = {'A_B_MDEA.pdf': {'folder': '/CVs/2026/August', 'url': 'old'}}
        in_rejected = {'A_B_MDEA.pdf': {'folder': '/CVs/2026/Rejected', 'url': 'old'}}

        remote = Mover()
        filed = _file_by_sheet(remote, in_month, 'rejected', received, dry_run=False)
        self.assertEqual(remote.moves, [('A_B_MDEA.pdf', '/CVs/2026/August', '/CVs/2026/Rejected')])
        self.assertEqual(filed['A_B_MDEA.pdf']['folder'], '/CVs/2026/Rejected')
        self.assertEqual(filed['A_B_MDEA.pdf']['url'], 'https://example.test/CVs/2026/Rejected/A_B_MDEA.pdf')

        remote = Mover()
        filed = _file_by_sheet(remote, in_rejected, 'main', received, dry_run=False)
        self.assertEqual(filed['A_B_MDEA.pdf']['folder'], '/CVs/2026/August',
                         'a candidate back on the main sheet has the CV returned to the month')

        for docs, sheet in ((in_month, 'main'), (in_rejected, 'rejected')):
            remote = Mover()
            self.assertEqual(_file_by_sheet(remote, docs, sheet, received, dry_run=False), docs)
            self.assertEqual(remote.moves, [], 'a CV already in the right folder is not touched')

        remote = Mover()
        self.assertEqual(_file_by_sheet(remote, in_month, 'rejected', received, dry_run=True), in_month)
        self.assertEqual(remote.moves, [], 'a dry run moves nothing')

        self.assertEqual(_file_by_sheet(Mover(fail=True), in_month, 'rejected', received, dry_run=False),
                         in_month, 'a failed move keeps the row pointing at the file it found')

    def test_rejected_result_only_patches_the_existing_p1_row(self):
        from hiring_agent.local_pipeline import _sync_row_to_sharepoint

        writes = []
        class Remote:
            table = 'HiringAgent_P1_Candidates'
            def list_rows(self):
                return [{'index': 0, 'values': dict(self_row)}]
            def row_values_at(self, index):
                return dict(self_row) if index == 0 else None
            def update_row(self, index, fields, current_values=None):
                writes.append((index, dict(fields)))
            def add_rejected_row(self, fields):
                raise AssertionError('P2 must not write P1 Rejected')
            def delete_row(self, index):
                raise AssertionError('P2 must not delete the P1 intake row')

        self_row = self.P1_ROW
        rejected = self._scored()
        rejected['Status'] = 'Rejected - Non-USA Location'
        _sync_row_to_sharepoint(Remote(), 'APP-1', 'rejected', rejected, self.P1_ROW)
        self.assertEqual(writes, [(0, {
            'Status': 'Rejected - Non-USA Location',
            'Resume Link': 'Marcus_Vance_AF6A',
        })])

    def test_a_scored_row_survives_the_next_poll(self):
        """The reset loop, end to end.

        P2 used to PATCH the whole row back, including 'Original Filename' and
        'Last Updated Date'. Both are hashed as input, so the next sync read the row as a
        NEW submission, reset it to 'New Email Received', cleared Retry Count and dropped
        the scored result - which P2 then scored again, rewriting the row every cycle.
        """
        import tempfile
        from hiring_agent.sqlite_store import SQLiteCandidateStore
        from hiring_agent.local_pipeline import _master_patch

        class Remote:
            def __init__(self, rows): self.rows = rows
            def list_rows(self): return [{'values': dict(v)} for v in self.rows]
            def list_rejected_rows(self): return []

        with patch.dict(os.environ,
                        {'HIRING_SQLITE_PATH': os.path.join(tempfile.mkdtemp(), 't.db')}):
            store = SQLiteCandidateStore()
            try:
                store.sync_from_client(Remote([self.P1_ROW]))
                store.save_by_id('APP-1', self._scored())
                self.assertEqual(store.all_rows('main')[0].values['Status'], 'Scored')

                # The master now holds P1's row plus only what P2 is allowed to write.
                master = dict(self.P1_ROW)
                master.update(_master_patch(self._scored()))
                store.sync_from_client(Remote([master]))

                survived = store.all_rows('main')[0].values
                self.assertEqual(survived['Status'], 'Scored',
                                 'a scored row must not be reset by the next poll')
                self.assertEqual(survived['Phone'], '(602) 555-0147',
                                 'the scored result must survive the next poll')
            finally:
                store.conn.close()

    def test_sync_from_client_refuses_shorter_snapshot_unless_forced(self):
        """A transiently short P1 read cannot silently remove a committed P2 candidate."""
        import tempfile
        from hiring_agent.sqlite_store import SQLiteCandidateStore

        class Remote:
            def __init__(self, rows): self.rows = rows
            def list_rows(self): return [{'values': dict(v)} for v in self.rows]
            def list_rejected_rows(self): return []

        row1 = dict(self.P1_ROW, **{'Application ID': 'APP-1'})
        row2 = dict(self.P1_ROW, **{'Application ID': 'APP-2'})
        with patch.dict(os.environ, {'HIRING_SQLITE_PATH': os.path.join(tempfile.mkdtemp(), 't.db')}):
            store = SQLiteCandidateStore()
            try:
                store.sync_from_client(Remote([row1, row2]))
                self.assertEqual(len(store.all_rows('main')), 2)

                with self.assertRaisesRegex(RuntimeError, 'APP-2'):
                    store.sync_from_client(Remote([row1]))
                self.assertEqual([r.app_id for r in store.all_rows('main')],
                                 ['APP-1', 'APP-2'])

                # Intentional administrative deletion remains available only when explicit.
                store.sync_from_client(Remote([row1]), force=True)
                active_rows = store.all_rows('main')
                self.assertEqual(len(active_rows), 1)
                self.assertEqual(active_rows[0].app_id, 'APP-1')
            finally:
                store.conn.close()

    def test_sync_from_client_ignores_p1s_historical_rejected_sheet(self):
        """P2's rejected state comes from SQLite/P2 master, not the P1 workbook."""
        import tempfile
        from hiring_agent.sqlite_store import SQLiteCandidateStore

        class Remote:
            def list_rows(self):
                return [{'values': dict(self_row)}]
            def list_rejected_rows(self):
                return [{'values': {
                    'Application ID': 'APP-OLD', 'Email': 'old@example.com',
                    'Status': 'Rejected - Non-USA Location',
                }}]

        self_row = self.P1_ROW
        with patch.dict(os.environ, {'HIRING_SQLITE_PATH': os.path.join(tempfile.mkdtemp(), 't.db')}):
            store = SQLiteCandidateStore()
            try:
                store.sync_from_client(Remote())
                self.assertIsNotNone(store.get('APP-1', 'main'))
                self.assertIsNone(store.get('APP-OLD', 'rejected'))
            finally:
                store.conn.close()

    def test_excel_candidate_store_guards_intake_table_against_profile_columns(self):
        """ExcelCandidateStore must never write profile fields to HiringAgent_P1_Candidates."""
        from hiring_agent.store import ExcelCandidateStore

        recorded = {}

        class MockClient:
            table = 'HiringAgent_P1_Candidates'
            def list_rows(self):
                return [{'index': 0, 'values': {'Application ID': 'APP-1', 'Status': 'New Email Received'}}]
            def update_row(self, index, fields, current_values=None):
                recorded.update(fields)

        sp_store = ExcelCandidateStore(MockClient())
        sp_store.save_by_id('APP-1', {
            'Status': 'Scored',
            'Resume Link': 'test.pdf',
            'Phone': '(555) 123-4567',
            'Location': 'San Jose, CA',
            'Category': 'AI/ML',
            'Suggested Role 1': 'AI Engineer',
        })
        self.assertEqual(recorded.get('Status'), 'Scored')
        self.assertEqual(recorded.get('Resume Link'), 'test.pdf')
        self.assertNotIn('Phone', recorded)
        self.assertNotIn('Location', recorded)
        self.assertNotIn('Category', recorded)
        self.assertNotIn('Suggested Role 1', recorded)

    def test_excel_store_keeps_rejected_outcome_on_p1_candidate_list(self):
        """The direct Excel fallback obeys the same P1/P2 ownership boundary."""
        from hiring_agent.store import ExcelCandidateStore

        calls = {'updates': [], 'adds': [], 'deletes': []}
        class MockClient:
            table = 'HiringAgent_P1_Candidates'
            def list_rows(self):
                return [{'index': 0, 'values': {
                    'Application ID': 'APP-1', 'Status': 'New Email Received'}}]
            def row_values_at(self, index):
                return self.list_rows()[0]['values'] if index == 0 else None
            def update_row(self, index, fields, current_values=None):
                calls['updates'].append(dict(fields))
            def add_rejected_row(self, fields):
                calls['adds'].append(dict(fields))
            def delete_row(self, index):
                calls['deletes'].append(index)

        moved = ExcelCandidateStore(MockClient()).move_to_rejected('APP-1', {
            'Application ID': 'APP-1',
            'Status': 'Rejected - Non-USA Location',
            'Resume Link': 'Candidate_ABCD.pdf',
            'Full Name': 'Extracted Name',
            'Country': 'India',
        })
        self.assertTrue(moved)
        self.assertEqual(calls['updates'], [{
            'Status': 'Rejected - Non-USA Location',
            'Resume Link': 'Candidate_ABCD.pdf',
        }])
        self.assertEqual(calls['adds'], [])
        self.assertEqual(calls['deletes'], [])


if __name__ == '__main__':
    unittest.main()
