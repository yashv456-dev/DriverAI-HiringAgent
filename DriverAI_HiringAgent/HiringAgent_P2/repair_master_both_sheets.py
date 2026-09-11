"""Master Workbook & SQLite Reconciliation Engine for Phase 2.

Repairs both sheets (CandidateList and Rejected) of the master file:
1. Purges the 17 duplicate ghost rows from CandidateList (authoritative copies exist on Rejected).
2. Purges the empty spacer row 101 from CandidateList.
3. Backfills recoverable missing fields (Phone, Full Name, Location, Education Dates, Years Exp)
   from verified resume texts using deterministic extraction with decimal-year preservation.
4. Preserves authentic Microsoft Excel Tables (HiringAgent_P1_Candidates ref A1:AG199,
   RejectedCandidates ref A1:AF80), cell styles, @ text formatting, and formulas.
5. Atomically updates SQLite candidates.db to match the pristine master state (198 main, 79 rejected).
6. Exports clean client report Candidate_List_Results.xlsx.
7. Backs up and uploads the repaired master workbook and client results to SharePoint.
"""
from __future__ import annotations

import argparse
import datetime
import json
import logging
import os
import re
import shutil
import sqlite3
import sys
from pathlib import Path

import openpyxl

APP_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(APP_DIR))

from dotenv import load_dotenv
load_dotenv(APP_DIR / ".env")

from hiring_agent.config import COLUMNS, REJECTED_COLUMNS
from hiring_agent.extraction import (
    _GAP_LITERALS,
    _extract_education_dates,
    _extract_experience,
    _extract_location,
    _extract_phone,
    clean_location_text,
    format_phone,
    location_is_plausible,
    sanitize_phone,
)
from hiring_agent.exporter import ExcelReportExporter
from hiring_agent.store import SQLiteCandidateStore

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s | %(levelname)-8s | %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S",
)
logger = logging.getLogger("repair_master")

SNAPSHOT_PATH = APP_DIR / "P2_Logs" / "audits" / "master_row_audit_20260906" / "master_snapshot.xlsx"
DOCS_PATH = APP_DIR / "P2_Logs" / "audits" / "master_row_audit_20260906" / "documents.json"
DB_PATH = APP_DIR / "candidates.db"

DUPLICATE_APP_IDS = {
    "APP-20260819-0140-WPOA",
    "APP-20260815-0302-WPDA",
    "APP-20260812-1003-WPXA",
    "APP-20260810-1751-WPEA",
    "APP-20260805-0132-WP7A",
    "APP-20260804-2313-WQGA",
    "APP-20260804-2208-WQMA",
    "APP-20260804-1923-WPKA",
    "APP-20260804-1910-WPPA",
    "APP-20260804-1557-WPXA",
    "APP-20260727-0731-WQOA",
    "APP-20260720-0313-WPSA",
    "APP-20260717-0400-WP5A",
    "APP-20260716-0957-WQJA",
    "APP-20260716-0805-WQPA",
    "APP-20260712-0546-WQKA",
    "APP-20260706-1100-WQTA",
}

KNOWN_OVERRIDE_PATCHES = {
    # Full Name missing on Scored candidate
    "APP-20260817-0617-WPWA": {
        "Full Name": "Charu Sneha Laguduva Ravi",
    },
    # Bare US phone numbers with NANP area codes
    "APP-20260706-1613-WQSA": {
        "Phone": "(623) 297-5055",
    },
    "APP-20260721-1330-WQZA": {
        "Phone": "(630) 806-4687",
    },
    # Missing location on confirmed US candidate
    "APP-20260602-1044-909E": {
        "Location": "Santa Clara, CA",
    },
    # Missing education dates
    "APP-20260503-1824-D4B8": {
        "Education Start Date": "Aug 2024",
        "Education End Date": "Sept 2025",
    },
    # Stated experience
    "APP-20260508-1438-96E1": {
        "Years Exp": "2 years",
    },
}


def is_gap(val: str | None) -> bool:
    if val is None:
        return True
    s = str(val).strip()
    return not s or s.lower() in _GAP_LITERALS


def load_documents_index() -> dict[str, str]:
    """Map filename or base name or app id to resume text."""
    if not DOCS_PATH.exists():
        logger.warning(f"Documents file not found at {DOCS_PATH}")
        return {}
    
    with open(DOCS_PATH, "r", encoding="utf-8") as f:
        docs = json.load(f)
    
    index = {}
    for doc in docs:
        name = doc.get("name", "").strip()
        path = doc.get("path", "").strip()
        text = doc.get("text", "")
        if name:
            index[name.lower()] = text
            index[Path(name).stem.lower()] = text
        if path:
            index[path.lower()] = text
            index[Path(path).name.lower()] = text
            index[Path(path).stem.lower()] = text
    return index


def find_resume_text(app_id: str, row_dict: dict, docs_index: dict[str, str]) -> str:
    """Locate resume text from documents index by multiple fallback keys."""
    # 1. By Application ID tail
    tail = app_id.split("-")[-1].lower()
    for k, v in docs_index.items():
        if tail in k and len(v) > 50:
            return v
    
    # 2. By Original Filename
    orig_fn = str(row_dict.get("Original Filename", "") or "").strip().lower()
    if orig_fn in docs_index:
        return docs_index[orig_fn]
    orig_stem = Path(orig_fn).stem.lower()
    if orig_stem in docs_index:
        return docs_index[orig_stem]
    
    # 3. By Resume URL filename
    url = str(row_dict.get("Resume URL", "") or "").strip()
    if url:
        url_fn = Path(url.split("?")[0]).name.lower()
        if url_fn in docs_index:
            return docs_index[url_fn]
        url_stem = Path(url_fn).stem.lower()
        if url_stem in docs_index:
            return docs_index[url_stem]
        
    # 4. By Full Name
    name = str(row_dict.get("Full Name", "") or "").strip().lower()
    if name and name not in ("missing", "not extracted"):
        name_parts = name.split()
        if len(name_parts) >= 2:
            for k, v in docs_index.items():
                if name_parts[0] in k and name_parts[-1] in k and len(v) > 50:
                    return v
    return ""


def repair_and_save_master_workbook(
    source_path: Path,
    output_path: Path,
    docs_index: dict[str, str],
) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    """Load original snapshot, perform in-place deletions and updates, preserving Excel XML tables."""
    logger.info(f"Modifying master snapshot in-place: {source_path}")
    wb = openpyxl.load_workbook(source_path)

    # ─────────────────────────────────────────────────────────────────
    # 1. CandidateList Sheet
    # ─────────────────────────────────────────────────────────────────
    ws_main = wb["CandidateList"]
    headers_main = [str(c.value or "").strip() for c in ws_main[1]]
    col_idx_map_main = {h: idx + 1 for idx, h in enumerate(headers_main) if h}

    # Identify rows to delete (empty row 101, unidentifiable rows, or 17 duplicates)
    rows_to_delete = []
    for r_idx in range(2, ws_main.max_row + 1):
        app_id = str(ws_main.cell(r_idx, 1).value or "").strip()
        if r_idx == 101 or not app_id or app_id in DUPLICATE_APP_IDS:
            rows_to_delete.append(r_idx)

    rows_to_delete.sort(reverse=True)
    logger.info(f"Purging {len(rows_to_delete)} stale/duplicate/empty rows from CandidateList...")
    for r_idx in rows_to_delete:
        ws_main.delete_rows(r_idx, 1)

    # Adjust table ref
    tab_main = ws_main.tables["HiringAgent_P1_Candidates"]
    tab_main.ref = f"A1:AG{ws_main.max_row}"
    logger.info(f"Updated CandidateList table ref: {tab_main.ref} ({ws_main.max_row - 1} candidate rows)")

    # Read, backfill, and update remaining CandidateList rows
    main_rows = []
    backfilled_main = {"Full Name": 0, "Phone": 0, "Location": 0, "Years Exp": 0, "Education Start Date": 0, "Education End Date": 0}

    for r_idx in range(2, ws_main.max_row + 1):
        row_dict = {h: str(ws_main.cell(r_idx, col_idx_map_main[h]).value or "").strip() for h in headers_main if h}
        app_id = row_dict.get("Application ID", "")
        if not app_id:
            continue

        # Check manual overrides
        if app_id in KNOWN_OVERRIDE_PATCHES:
            for k, v in KNOWN_OVERRIDE_PATCHES[app_id].items():
                if is_gap(row_dict.get(k)):
                    row_dict[k] = v
                    ws_main.cell(r_idx, col_idx_map_main[k], v)
                    if k in backfilled_main:
                        backfilled_main[k] += 1
                    logger.info(f"   CandidateList override for {app_id}: {k} = '{v}'")

        # Resume text backfill
        txt = find_resume_text(app_id, row_dict, docs_index)
        if txt:
            if is_gap(row_dict.get("Years Exp")):
                exp = _extract_experience(txt)
                if exp:
                    row_dict["Years Exp"] = exp
                    ws_main.cell(r_idx, col_idx_map_main["Years Exp"], exp)
                    backfilled_main["Years Exp"] += 1
            if is_gap(row_dict.get("Education Start Date")) or is_gap(row_dict.get("Education End Date")):
                s_dt, e_dt = _extract_education_dates(txt)
                if is_gap(row_dict.get("Education Start Date")) and not is_gap(s_dt):
                    row_dict["Education Start Date"] = s_dt
                    ws_main.cell(r_idx, col_idx_map_main["Education Start Date"], s_dt)
                    backfilled_main["Education Start Date"] += 1
                if is_gap(row_dict.get("Education End Date")) and not is_gap(e_dt):
                    row_dict["Education End Date"] = e_dt
                    ws_main.cell(r_idx, col_idx_map_main["Education End Date"], e_dt)
                    backfilled_main["Education End Date"] += 1
            if is_gap(row_dict.get("Phone")):
                raw_p = _extract_phone(txt)
                if raw_p:
                    clean_p = sanitize_phone(format_phone(raw_p))
                    if not is_gap(clean_p):
                        row_dict["Phone"] = clean_p
                        ws_main.cell(r_idx, col_idx_map_main["Phone"], clean_p)
                        backfilled_main["Phone"] += 1
            if is_gap(row_dict.get("Location")):
                loc = clean_location_text(_extract_location(txt) or "")
                if loc and location_is_plausible(loc):
                    row_dict["Location"] = loc
                    ws_main.cell(r_idx, col_idx_map_main["Location"], loc)
                    backfilled_main["Location"] += 1

        # Text format for phone and dates
        for col_name in ("Phone", "Education Start Date", "Education End Date"):
            if col_name in col_idx_map_main:
                ws_main.cell(r_idx, col_idx_map_main[col_name]).number_format = "@"

        main_rows.append(row_dict)

    logger.info(f"[CandidateList] Backfill counts: {backfilled_main}")

    # ─────────────────────────────────────────────────────────────────
    # 2. Rejected Sheet
    # ─────────────────────────────────────────────────────────────────
    ws_rej = wb["Rejected"]
    headers_rej = [str(c.value or "").strip() for c in ws_rej[1]]
    col_idx_map_rej = {h: idx + 1 for idx, h in enumerate(headers_rej) if h}

    tab_rej = ws_rej.tables["RejectedCandidates"]
    tab_rej.ref = f"A1:AF{ws_rej.max_row}"
    logger.info(f"Rejected table ref: {tab_rej.ref} ({ws_rej.max_row - 1} candidate rows)")

    rejected_rows = []
    backfilled_rej = {"Years Exp": 0, "Phone": 0, "Education Start Date": 0, "Education End Date": 0}

    for r_idx in range(2, ws_rej.max_row + 1):
        row_dict = {h: str(ws_rej.cell(r_idx, col_idx_map_rej[h]).value or "").strip() for h in headers_rej if h}
        app_id = row_dict.get("Application ID", "")
        if not app_id:
            continue

        txt = find_resume_text(app_id, row_dict, docs_index)
        if txt:
            if is_gap(row_dict.get("Years Exp")):
                exp = _extract_experience(txt)
                if exp:
                    row_dict["Years Exp"] = exp
                    ws_rej.cell(r_idx, col_idx_map_rej["Years Exp"], exp)
                    backfilled_rej["Years Exp"] += 1
            if is_gap(row_dict.get("Education Start Date")) or is_gap(row_dict.get("Education End Date")):
                s_dt, e_dt = _extract_education_dates(txt)
                if is_gap(row_dict.get("Education Start Date")) and not is_gap(s_dt):
                    row_dict["Education Start Date"] = s_dt
                    ws_rej.cell(r_idx, col_idx_map_rej["Education Start Date"], s_dt)
                    backfilled_rej["Education Start Date"] += 1
                if is_gap(row_dict.get("Education End Date")) and not is_gap(e_dt):
                    row_dict["Education End Date"] = e_dt
                    ws_rej.cell(r_idx, col_idx_map_rej["Education End Date"], e_dt)
                    backfilled_rej["Education End Date"] += 1
            if is_gap(row_dict.get("Phone")):
                raw_p = _extract_phone(txt)
                if raw_p:
                    clean_p = sanitize_phone(format_phone(raw_p))
                    if not is_gap(clean_p):
                        row_dict["Phone"] = clean_p
                        ws_rej.cell(r_idx, col_idx_map_rej["Phone"], clean_p)
                        backfilled_rej["Phone"] += 1

        for col_name in ("Phone", "Education Start Date", "Education End Date"):
            if col_name in col_idx_map_rej:
                ws_rej.cell(r_idx, col_idx_map_rej[col_name]).number_format = "@"

        rejected_rows.append(row_dict)

    logger.info(f"[Rejected] Backfill counts: {backfilled_rej}")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    logger.info(f"[OK] Pristine repaired master file written to: {output_path}")

    return main_rows, rejected_rows


def sync_sqlite_database(
    main_rows: list[dict[str, str]],
    rejected_rows: list[dict[str, str]],
    db_path: Path = DB_PATH,
) -> None:
    """Sync candidates.db with repaired and deduplicated candidate rows."""
    logger.info(f"Syncing SQLite database at {db_path}...")
    conn = sqlite3.connect(db_path, timeout=30.0)
    conn.execute("PRAGMA journal_mode = WAL;")
    conn.execute("PRAGMA synchronous = NORMAL;")

    with conn:
        conn.execute("DELETE FROM candidates;")
        for r in main_rows:
            app_id = r.get("Application ID", "").strip()
            status = r.get("Status", "").strip()
            val_json = json.dumps(r, ensure_ascii=False)
            conn.execute(
                "INSERT INTO candidates (app_id, sheet, status, values_json, version) VALUES (?, 'main', ?, ?, 1)",
                (app_id, status, val_json),
            )

        for r in rejected_rows:
            app_id = r.get("Application ID", "").strip()
            status = r.get("Status", "").strip()
            val_json = json.dumps(r, ensure_ascii=False)
            conn.execute(
                "INSERT INTO candidates (app_id, sheet, status, values_json, version) VALUES (?, 'rejected', ?, ?, 1)",
                (app_id, status, val_json),
            )

    cur = conn.cursor()
    cur.execute("SELECT sheet, count(*) FROM candidates GROUP BY sheet")
    counts = dict(cur.fetchall())
    logger.info(f"[OK] Database sync complete. Counts: {counts}")

    cur.execute(
        """
        SELECT app_id FROM candidates WHERE sheet = 'main'
        INTERSECT
        SELECT app_id FROM candidates WHERE sheet = 'rejected'
        """
    )
    overlap = cur.fetchall()
    if overlap:
        logger.error(f"[FAIL] Unexpected cross-sheet overlap in DB: {overlap}")
    else:
        logger.info("[PASS] Zero cross-sheet duplicates in SQLite database!")
    conn.close()


def run_repair(upload_to_sharepoint: bool = False):
    """Execute complete master file repair and synchronization."""
    ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    logger.info("=" * 70)
    logger.info(f"STARTING MASTER WORKBOOK & DB REPAIR ({ts})")
    logger.info("=" * 70)

    docs_index = load_documents_index()
    logger.info(f"Loaded {len(docs_index)} document entries for resume text backfill.")

    out_master = APP_DIR / "P2_Final_Results" / "Sharepoint_Master_File.xlsx"
    out_backup = APP_DIR / "P2_Final_Results" / f"Sharepoint_Master_File_{ts}.xlsx"

    clean_main, clean_rej = repair_and_save_master_workbook(
        source_path=SNAPSHOT_PATH,
        output_path=out_master,
        docs_index=docs_index,
    )
    shutil.copy2(out_master, out_backup)
    logger.info(f"[OK] Backup copy saved at: {out_backup}")

    assert len(clean_main) == 198, f"Expected 198 CandidateList rows, got {len(clean_main)}"
    assert len(clean_rej) == 79, f"Expected 79 Rejected rows, got {len(clean_rej)}"
    main_ids = {r["Application ID"] for r in clean_main}
    rej_ids = {r["Application ID"] for r in clean_rej}
    assert not (main_ids & rej_ids), f"Cross-sheet overlap found: {main_ids & rej_ids}"
    logger.info("[PASS] Count and deduplication validation passed (198 main, 79 rejected, 0 overlap).")

    # Sync SQLite database
    sync_sqlite_database(clean_main, clean_rej, DB_PATH)

    # Generate fresh Candidate_List_Results.xlsx
    store = SQLiteCandidateStore(DB_PATH)
    exporter = ExcelReportExporter(store)
    client_results_path = APP_DIR / "P2_Final_Results" / "Candidate_List_Results.xlsx"
    exporter.generate_workbook(client_results_path)
    logger.info(f"[OK] Generated client results report at: {client_results_path}")

    # Upload to SharePoint
    if upload_to_sharepoint:
        logger.info("Connecting to SharePoint to upload repaired master file...")
        from sharepoint_client import SharePointClient
        client = SharePointClient()

        # Remote backup of previous master file
        try:
            old_bytes = client.download_file("/Master_Files", "Sharepoint_Master_File.xlsx")
            backup_name = f"Sharepoint_Master_File_backup_{ts}.xlsx"
            client.upload_file("/Master_Files", backup_name, old_bytes)
            logger.info(f"[OK] SharePoint remote backup created: /Master_Files/{backup_name}")
        except Exception as e:
            logger.warning(f"Could not create remote SharePoint backup: {e}")

        # Upload repaired master file
        with open(out_master, "rb") as f:
            master_bytes = f.read()
        client.upload_file("/Master_Files", "Sharepoint_Master_File.xlsx", master_bytes)
        logger.info("[OK] Uploaded repaired Sharepoint_Master_File.xlsx to SharePoint (/Master_Files)")

        # Upload client results
        with open(client_results_path, "rb") as f:
            client_bytes = f.read()
        client.upload_file("/", "Candidate_List_Results.xlsx", client_bytes)
        logger.info("[OK] Uploaded Candidate_List_Results.xlsx to SharePoint root (/Candidate_List_Results.xlsx)")

        # Live verification via Graph API
        logger.info("Performing live post-upload verification on SharePoint via Graph API...")
        client._columns_cache.clear()
        live_main = client.list_rows()
        live_rej = client.list_rejected_rows()
        logger.info(f"Live SharePoint table counts: {len(live_main)} CandidateList, {len(live_rej)} Rejected.")
        live_main_ids = {r["values"].get("Application ID") for r in live_main}
        live_rej_ids = {r["values"].get("Application ID") for r in live_rej}
        live_inter = live_main_ids & live_rej_ids
        if live_inter:
            logger.error(f"[FAIL] Live SharePoint still has overlap: {live_inter}")
        else:
            logger.info("[PASS] Live SharePoint verified: exactly 0 cross-sheet duplicates!")

    logger.info("=" * 70)
    logger.info("MASTER WORKBOOK & DB REPAIR SUCCESSFULLY COMPLETED")
    logger.info("=" * 70)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Repair both sheets of master workbook")
    parser.add_argument("--upload", action="store_true", help="Upload repaired master file to SharePoint")
    args = parser.parse_args()

    run_repair(upload_to_sharepoint=args.upload)

