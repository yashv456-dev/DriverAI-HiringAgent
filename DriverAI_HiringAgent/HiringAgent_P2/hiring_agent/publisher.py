"""Generate reports from committed state and retry delivery without rerunning inference."""
from __future__ import annotations

import hashlib
import io
import json
import os
import datetime as dt
from pathlib import Path
from urllib.parse import unquote

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

from . import config as cfg
from .excel_output import finish_sheet
from .sqlite_store import BASE


def safe_target(target, intake):
    def canonical(value):
        parts = unquote(str(value)).replace('\\', '/').strip('/').split('/')
        if any(p in ('.', '..') for p in parts):
            raise ValueError('Report path must not contain traversal segments')
        return '/'.join(parts).casefold()
    if canonical(target) == canonical(intake):
        raise ValueError('Refusing to replace the P1 intake workbook with a generated report')


def _workbook_shape(data):
    """(sheet title, row count, column count) per sheet, for an in-memory xlsx."""
    import io
    wb = load_workbook(io.BytesIO(data), read_only=True)
    try:
        return [(ws.title, ws.max_row, ws.max_column) for ws in wb.worksheets]
    finally:
        wb.close()


def _cell_value(value):
    """Stable workbook value for semantic post-upload comparison."""
    if value is None:
        return ""
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    return value


def _workbook_content(data):
    """Content that must survive SharePoint's server-side xlsx rewrite exactly.

    Zip bytes and workbook metadata unrelated to the report can change on upload. Candidate
    identities, every displayed value, hyperlink target, sheet/table layout, and committed
    generation cannot.  Comparing this representation catches a same-sized but stale or
    cross-wired workbook, which the former row/column-count check accepted.
    """
    wb = load_workbook(io.BytesIO(data), data_only=False, read_only=False)
    try:
        sheets = []
        for ws in wb.worksheets:
            rows = []
            links = []
            for row in ws.iter_rows(min_row=1, max_row=ws.max_row,
                                    min_col=1, max_col=ws.max_column):
                rows.append([_cell_value(cell.value) for cell in row])
                links.append([str(cell.hyperlink.target or "") if cell.hyperlink else ""
                              for cell in row])
            sheets.append({
                "title": ws.title,
                "rows": rows,
                "hyperlinks": links,
                "tables": sorted((name, ws.tables[name].ref) for name in ws.tables.keys()),
            })
        return {
            "generation": str(wb.properties.description or ""),
            "sheets": sheets,
        }
    finally:
        wb.close()


def _verify_report(remote, folder, name, data):
    """Confirm the stored report is the report we sent, by CONTENT not by bytes.

    SharePoint rewrites an uploaded .xlsx server-side - the same workbook came back 164,931
    bytes against the 157,156 uploaded (2026-09-08), both valid archives - so a sha256 of the
    raw bytes can never match and the publication aborted after every successful upload. What
    has to hold is that the stored file is a readable workbook with the sheets, rows and
    columns we published; re-zipping does not change any of those.
    """
    stored = remote.download_file(folder, name)
    if stored == data:
        return _workbook_content(stored)
    try:
        actual = _workbook_content(stored)
    except Exception as error:
        raise ValueError(f'Stored report is not a readable workbook: {name} ({error})')
    expected = _workbook_content(data)
    if actual != expected:
        def ids(content):
            out = []
            for sheet in content.get('sheets', []):
                rows = sheet.get('rows', [])
                if not rows or 'Application ID' not in rows[0]:
                    continue
                col = rows[0].index('Application ID')
                out.extend(str(row[col]) for row in rows[1:]
                           if col < len(row) and str(row[col]).strip())
            return out
        expected_ids, actual_ids = ids(expected), ids(actual)
        raise ValueError(
            f'Stored report content does not match what was published: {name}; '
            f'expected generation={expected["generation"]!r}, '
            f'actual generation={actual["generation"]!r}, '
            f'missing IDs={sorted(set(expected_ids) - set(actual_ids))}, '
            f'extra IDs={sorted(set(actual_ids) - set(expected_ids))}'
        )
    return actual


def build_master(store, path, generation):
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.description = f'P2 committed generation {generation}'
    def _unfamiliar(values):
        """Read against the whole catalogue and matched nothing in it.

        Keyed on Status, not on Category being 'General': a genuinely scored candidate can
        land in General when their role maps to no known category, and those people belong
        on CandidateList with everyone else.
        """
        return (str(values.get('Status', '') or '').strip()
                == cfg.STATUS_NO_MATCHING_ROLE)

    seen = set()
    expected_by_sheet = {}
    # 'main' is written twice, split by _unfamiliar into complementary halves - so the
    # duplicate guard below still holds (every application lands on exactly one sheet) and
    # no candidate is listed in two places. Rejected is the non-USA outcome and is separate
    # from having no opening to match, which is why this is a third sheet and not that one.
    for sheet, title, columns, table_name, keep in (
            ('main', 'CandidateList', cfg.COLUMNS, 'HiringAgent_P1_Candidates',
             lambda values: not _unfamiliar(values)),
            ('rejected', 'Rejected', cfg.REJECTED_COLUMNS, 'RejectedCandidates',
             lambda values: True),
            ('main', 'Unfamiliar Role List', cfg.COLUMNS, 'UnfamiliarRoleCandidates',
             _unfamiliar)):
        ws = wb.create_sheet(title)
        ws.append(columns)
        # header / blank spacer / first candidate, matching P1's intake workbook and the
        # client sheet so all three read the same way. Carries no Application ID, so every
        # reader already skips it.
        ws.append([''] * len(columns))
        expected_by_sheet[title] = []
        for row in store.all_rows(sheet):
            if not row.app_id:
                continue
            if not keep(row.values):
                continue
            if row.app_id in seen:
                raise ValueError(f'Ambiguous active application {row.app_id}; publication deferred')
            seen.add(row.app_id)
            expected_by_sheet[title].append(row.app_id)
            ws.append([row.values.get(col, '') for col in columns])
            for cell in ws[ws.max_row]:
                # Resume/mail values are literal strings, never executable formulas.
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                cell.number_format = '@'
        table = Table(displayName=table_name, ref=f'A1:{get_column_letter(len(columns))}{ws.max_row}')
        # showRowStripes drops with the doubt fill: banding is row colour too, and the
        # brief was that no row carries a colour. The header band stays, so the table
        # still reads as a table.
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium9', showRowStripes=False)
        finish_sheet(ws, columns, table=table)
        ws.add_table(table)
        for i, col in enumerate(columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = min(max(len(col) + 3, 18), 42)
        if 'Resume URL' in columns and 'Resume Link' in columns:
            url_col, link_col = columns.index('Resume URL') + 1, columns.index('Resume Link') + 1
            for row in range(2, ws.max_row + 1):
                url = str(ws.cell(row, url_col).value or '')
                if url.startswith(('https://', 'http://')):
                    from urllib.parse import unquote, urlparse, parse_qs
                    qs = parse_qs(urlparse(url).query)
                    if "file" in qs and qs["file"]:
                        display = qs["file"][0]
                    else:
                        display = unquote(url.rstrip("/").rsplit("/", 1)[-1].split("?")[0])
                    cell = ws.cell(row, link_col, display or 'Resume')
                    cell.hyperlink = url
                    # The loop above stamps every cell as Text ('@') so a stored value can
                    # never be evaluated as a formula. Correct for data, wrong here: it left
                    # the link rendering as ordinary black text with no underline, so it did
                    # not READ as a link even though clicking it worked. Restore the default
                    # format and apply Excel's own Hyperlink styling to this cell only.
                    cell.number_format = 'General'
                    cell.font = Font(color='0563C1', underline='single')
    temp = Path(path).with_suffix('.tmp.xlsx')
    wb.save(temp)
    wb.close()
    check = load_workbook(temp)
    try:
        # The Unfamiliar sheet is always written, empty or not, so this stays an exact match
        # rather than a subset test - a missing sheet is still a schema failure.
        if (check.sheetnames != ['CandidateList', 'Rejected', 'Unfamiliar Role List']
                or any(not ws.tables for ws in check)):
            raise ValueError('Master workbook schema validation failed')
        # Application ID need not remain column A forever.
        expected = len(seen)
        id_counts = 0
        for ws in check:
            columns = [c.value for c in ws[1]]
            id_col = columns.index('Application ID') + 1
            actual_ids = [str(ws.cell(i, id_col).value) for i in range(2, ws.max_row + 1)
                          if ws.cell(i, id_col).value]
            id_counts += len(actual_ids)
            if actual_ids != expected_by_sheet[ws.title]:
                raise ValueError(
                    f'Master workbook Application ID validation failed for {ws.title}: '
                    f'expected {expected_by_sheet[ws.title]}, got {actual_ids}'
                )
        if id_counts != expected:
            raise ValueError('Master workbook count validation failed')
    finally:
        check.close()
    temp.replace(path)


def publish_results(store, remote, *, upload=True, output_dir=None, force=False):
    from .exporter import ExcelReportExporter
    from sharepoint_client import SharePointError
    output = Path(output_dir or os.getenv('HIRING_PUBLISH_DIR') or BASE / 'P2_Final_Results')
    output.mkdir(parents=True, exist_ok=True)
    record = store.conn.execute('SELECT * FROM publication WHERE id=1').fetchone()
    revision = record['revision']
    master_target = os.getenv('HIRING_MASTER_REPORT_PATH', '/Master_Files/P2-MasterFile.xlsx')
    client_target = os.getenv('HIRING_CLIENT_REPORT_PATH', '/Candidate_List_Results.xlsx')
    for target in (master_target, client_target):
        safe_target(target, remote._wb_path)
    if master_target.strip('/').casefold() == client_target.strip('/').casefold():
        raise ValueError('Master and client results must have distinct targets')
    generation = output / f'generation_{revision}'
    generation.mkdir(exist_ok=True)
    master = generation / 'P2-MasterFile.xlsx'
    client = generation / 'Candidate_List_Results.xlsx'
    if not force and upload and record['published_revision'] == revision and master.exists() and client.exists():
        try:
            client_folder, _, client_name = client_target.strip('/').rpartition('/')
            _verify_report(remote, client_folder, client_name, client.read_bytes())
            master_folder, _, master_name = master_target.strip('/').rpartition('/')
            _verify_report(remote, master_folder, master_name, master.read_bytes())
            if record['error']:
                with store.conn:
                    store.conn.execute('UPDATE publication SET error=? WHERE id=1', ('',))
            return client
        except Exception as error:
            cfg.logger.warning('Published revision %s failed semantic verification; '
                               'rebuilding and republishing: %s', revision, error)
    # Always rebuild master and client fresh from store so published files are never stale.
    build_master(store, master, revision)
    temp_client = client.with_suffix('.tmp.xlsx')
    ExcelReportExporter(store).generate_workbook(temp_client)
    wb = load_workbook(temp_client)
    wb.properties.description = f'P2 committed generation {revision}'
    for ws in wb:
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
    wb.save(temp_client)
    wb.close()
    temp_client.replace(client)
    manifest = {'generation': revision, 'files': [],
                'import_conflicts': [dict(r) for r in store.conn.execute('SELECT app_id,reason FROM import_conflicts')]}
    if not upload:
        return client
    targets = ((client, client_target), (master, master_target))
    previous = {}
    uploaded = []
    try:
        # Capture the currently published pair before changing either file. A successful
        # first upload can then be undone if the second workbook is locked or fails its
        # read-back check.
        for _path, target in targets:
            folder, _, name = target.strip('/').rpartition('/')
            try:
                previous[(folder, name)] = remote.download_file(folder, name)
            except SharePointError as error:
                if error.status_code != 404:
                    raise
                previous[(folder, name)] = None

        # EXACTLY three workbooks exist on the tenant: P1's intake master, this master report,
        # and this client sheet - and P2 only ever writes the last two, in place. Version
        # history lives in the local generation_<N> folders, which is why no '.r<N>' sibling
        # and no remote pointer file are uploaded: every extra remote copy was one more
        # 'master file' for a human to mistake for the real one. `published_revision` in the
        # local publication table already answers 'what is live', so the pointer was duplicate
        # state that could itself go stale.
        # Publish the client sheet first. It is the file most often held open in Excel; if
        # that write is locked, P2 Master remains untouched. If the second write fails, the
        # first is restored from `previous` below, keeping both remote files on one revision.
        for path, target in targets:
            folder, _, name = target.strip('/').rpartition('/')
            data = path.read_bytes()
            remote.upload_file(folder, name, data)
            uploaded.append((folder, name))
            _verify_report(remote, folder, name, data)
            manifest['files'].append({'path': '/'.join((folder, name)),
                                      'sha256': hashlib.sha256(data).hexdigest()})
        (generation / 'published.json').write_text(json.dumps(manifest, indent=2), encoding='utf-8')
        with store.conn:
            store.conn.execute('UPDATE publication SET published_revision=?,error=? WHERE id=1', (revision, ''))
    except Exception as error:
        rollback_errors = []
        for folder, name in reversed(uploaded):
            old = previous[(folder, name)]
            try:
                if old is None:
                    remote.delete_file(folder, name)
                else:
                    remote.upload_file(folder, name, old)
                    _verify_report(remote, folder, name, old)
            except Exception as rollback_error:
                rollback_errors.append(f'{folder}/{name}: {rollback_error}')
        detail = str(error)
        if uploaded and not rollback_errors:
            detail += ' (partial publication rolled back successfully)'
        elif rollback_errors:
            detail += ' (ROLLBACK FAILED: ' + '; '.join(rollback_errors) + ')'
        with store.conn:
            store.conn.execute('UPDATE publication SET error=? WHERE id=1', (detail,))
        cfg.logger.warning('Results committed locally; publication pending: %s', detail)
    return client
