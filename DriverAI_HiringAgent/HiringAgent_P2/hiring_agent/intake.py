"""Local intake: score a folder of resumes into the local Excel workbook.

Behaviour: for each PDF/DOCX, extract text -> read details (Ollama brain or offline) ->
score role fit (Ollama brain or keyword) -> apply the USA geo filter -> ADD a new row or
UPDATE an existing one (matched by resume filename, then email). New resumes are filed into
a dated <Year>/<Month>/ folder. Rejected (non-USA) rows go to a 'Rejected' sheet.

Also exposes score_existing_excel(): run the scoring brain over the rows of an arbitrary
workbook in place.
"""
import datetime
import itertools
import re
import time
from pathlib import Path

import pandas as pd

import hiring_agent.config as _cfg
from hiring_agent.config import (
    COLUMNS, STATUS_LOCATION_REVIEW, STATUS_SCORED,
    GEO_FILTER_USA_ONLY, logger,
)
from hiring_agent.extraction import (
    extract_text_from_bytes, extract_candidate_details_smart, resolve_full_name,
    infer_looking_for_role, infer_missing_portfolios, _GAP_LITERALS,
    sanitize_phone, format_phone,
)
from hiring_agent.scoring import suggested_roles, assign_category
from hiring_agent.jd_sources import get_active_roles
from hiring_agent.geo import GeoDecision, classify_location_usa
from hiring_agent.ollama_scorer import ollama_health

_APP_SEQ = itertools.count(1)
_PREFIX_RE = re.compile(r"^APP-\d{8}-\d{6}_", re.IGNORECASE)


def _new_app_id() -> str:
    # LOCAL, not UTC (2026-08-27). The row this ref is minted for stores its
    # "Received Date" from datetime.now() - local wall-clock - and P1 now mints its own
    # APP- refs from the local receivedDateTime too. Left on UTC, a 23:30 local upload
    # got a ref dated the NEXT day, disagreeing with its own row and with every other
    # reference in the system (Arizona is UTC-7).
    return f"APP-{datetime.datetime.now():%Y%m%d-%H%M%S}-{next(_APP_SEQ):02d}"


def _strip_prefix(name: str) -> str:
    """'APP-20260619-101010_jane_cv.pdf' -> 'jane_cv.pdf' (lowercased)."""
    return _PREFIX_RE.sub("", (name or "").strip()).lower()


def _row_basenames(cell) -> set:
    """The original resume name(s) in a row's 'Original Filename' cell, prefix-stripped."""
    return {_strip_prefix(p) for p in str(cell or "").split(",") if p.strip()}


def _matches(name: str, email: str, rf_cell, email_cell) -> bool:
    base = _strip_prefix(name)
    if base and base in _row_basenames(rf_cell):
        return True
    e = (email or "").strip().lower()
    return bool(e) and e != "not extracted" and e == str(email_cell or "").strip().lower()


def _save_local(df, target) -> None:
    """Write the workbook with a locked-file rescue (Excel open during a run)."""
    df = df.reindex(columns=COLUMNS)
    try:
        df.to_excel(target, index=False)
    except PermissionError:
        fb = target.with_name(f"{target.stem}_locked_{datetime.datetime.now():%Y%m%d-%H%M%S}.xlsx")
        df.to_excel(fb, index=False)
        logger.warning(f"   '{target.name}' is locked - wrote to '{fb.name}' instead.")


def _save_rejected_local(rows: list, target) -> None:
    """Append rejected rows to a 'Rejected' sheet in the same workbook."""
    from openpyxl import load_workbook
    df_rej = pd.DataFrame(rows, columns=COLUMNS)
    try:
        wb = load_workbook(target)
        if "Rejected" in wb.sheetnames:
            ws = wb["Rejected"]
        else:
            ws = wb.create_sheet("Rejected")
            ws.append(COLUMNS)
        for _, row in df_rej.iterrows():
            ws.append([row.get(c, "") for c in COLUMNS])
        wb.save(target)
        logger.info(f"   {len(rows)} rejected row(s) written to 'Rejected' sheet in {target.name}")
    except Exception as e:
        fb = target.with_name(f"{target.stem}_rejected_{datetime.datetime.now():%Y%m%d-%H%M%S}.xlsx")
        df_rej.to_excel(fb, index=False, sheet_name="Rejected")
        logger.warning(f"   could not write Rejected sheet ({e}) - saved to '{fb.name}'")


# Column name -> extraction-dict key, for the Step A/E snapshot-and-diff log below.
# Mirrors sharepoint_scoring.py's _P2_COL_TO_KEY so --score-folder and --score-sharepoint
# give identical column-level log visibility.
_P2_COL_TO_KEY = {
    "Full Name": "full_name", "Phone": "phone", "Location": "location",
    "Country": "country", "Current Skills": "skills",
    "Looking For Role": "looking_for_role",
    "Portfolio 1": "portfolio_1", "Portfolio 2": "portfolio_2", "Portfolio 3": "portfolio_3",
}


def _cell(v) -> str:
    """Excel-safe scalar-to-string. A bare `v or ""` breaks on pandas NaN (NaN is truthy),
    so blank Excel cells would otherwise log as the literal string 'nan'."""
    if v is None or (isinstance(v, float) and v != v):   # `v != v` is True only for NaN
        return ""
    return str(v).strip()


def _is_gap(v) -> bool:
    return _cell(v).lower() in _GAP_LITERALS


def _log_column_diff(before: dict, final: dict, after_extract: dict) -> None:
    """Step E: diff Step A's pre-extraction snapshot against the final row state, per
    column - mirrors sharepoint_scoring.py's _log_column_diff."""
    unchanged, filled, healed_cols = [], [], []
    method_lines = []
    for col in COLUMNS:
        was, now = _cell(before.get(col)), _cell(final.get(col))
        was_gap, now_gap = _is_gap(was), _is_gap(now)
        if was_gap and not now_gap:
            filled.append(col)
        elif not was_gap and now != was:
            healed_cols.append(col)
        else:
            unchanged.append(col)

        if col not in _P2_COL_TO_KEY:
            continue
        if was_gap and not now_gap:
            mid = _cell(after_extract.get(col))
            method = ("deterministic-only / AI-confirmed hint" if not _is_gap(mid)
                      else "AI-inferred-from-empty")
            method_lines.append(f"{col}: {method}")
        elif not was_gap and now != was:
            # Don't assert "AI" when it might just be a fresh regex re-scan - after_extract
            # tells "changed during extraction" apart from "changed later" without
            # overclaiming AI involvement (mirrors sharepoint_scoring.py's same fix).
            mid = _cell(after_extract.get(col))
            method = ("corrected during extraction (regex re-scan or AI-confirm)" if mid != was
                      else "corrected later")
            method_lines.append(f"{col}: {method}")

    logger.info(f"   Diff    : {len(unchanged)} unchanged | {len(filled)} filled | "
                f"{len(healed_cols)} healed")
    if method_lines:
        logger.info("   Method  : " + "; ".join(method_lines))


def run_intake(path: str = None, dry_run: bool = False, excel_file=None) -> dict:
    """Score every PDF/DOCX in a local folder and upsert into the local workbook.

    Returns a summary dict {added, updated, rejected, skipped, errors, workbook}.
    """
    roles = get_active_roles()
    folder = Path(path or "")
    if not folder.is_dir():
        logger.error(f"Not a folder: {path}")
        return {"error": f"Not a folder: {path}"}

    target_xlsx = Path(excel_file) if excel_file else _cfg.active_excel_file()
    df = (pd.read_excel(target_xlsx) if Path(target_xlsx).exists()
          else pd.DataFrame(columns=COLUMNS)).reindex(columns=COLUMNS)
    # Force object dtype so per-cell updates of mixed text/number fields never hit a
    # numpy dtype clash (an all-empty column reads back from Excel as all-NaN float64,
    # which pandas 3.x refuses to accept a string into).
    df = df.astype(object)

    resumes = [(f.name, f.read_bytes()) for f in sorted(folder.iterdir())
               if f.is_file() and f.suffix.lower() in (".pdf", ".docx")]
    logger.info(f"Intake | folder='{folder}' | workbook='{target_xlsx.name}' | "
                f"{len(resumes)} file(s) | {'DRY-RUN' if dry_run else 'LIVE'} | "
                f"geo_filter={'USA-only' if GEO_FILTER_USA_ONLY else 'off'}")
    _ok, _brain = ollama_health()
    (logger.info if _ok else logger.warning)(f"   {_brain}")

    added = updated = skipped = errors = rejected = 0
    now = datetime.datetime.now()
    rejected_rows = []

    for name, raw in resumes:
        try:
            t0 = time.perf_counter()
            text = extract_text_from_bytes(raw, name)
            if not text.strip():
                logger.warning(f"   no readable text in '{name}' (scanned/image PDF?) - skipped.")
                skipped += 1
                continue
            from hiring_agent.extraction import EXTRACTION_SOURCE
            d = extract_candidate_details_smart(text)
            if getattr(_cfg, "REQUIRE_AI", False) and EXTRACTION_SOURCE.get("value") != "ollama":
                logger.warning(f"   LLM extraction unavailable for '{name}' ({EXTRACTION_SOURCE.get('value')}); "
                               "skipped under REQUIRE_AI (regex fallback disabled).")
                skipped += 1
                continue

            # Snapshot the Tier 1/2 fields right after extraction, before the portfolio
            # gap-fill or role-inference below can still touch them - Step E uses this to
            # tell "filled during extraction" apart from "filled later".
            after_extract_snapshot = {col: d.get(key, "") for col, key in _P2_COL_TO_KEY.items()}

            skills = d.get("skills", "Not extracted")
            role_pref = d.get("looking_for_role", "Not extracted")
            if role_pref.strip().lower() in _GAP_LITERALS:
                # No mail body in local-folder intake - infer from resume skills only.
                role_pref = infer_looking_for_role(text, "", skills)
            res = suggested_roles(skills, role_pref, roles=roles, resume_text=text)
            if getattr(_cfg, "REQUIRE_AI", False) and res.get("source") != "ollama":
                logger.warning(f"   LLM role scoring unavailable for '{name}' ({res.get('source')}); "
                               "skipped under REQUIRE_AI (keyword fallback disabled).")
                skipped += 1
                continue
            r1, r2, r3 = res["role_1"], res["role_2"], res.get("role_3", "")
            category = assign_category(r1, skills, role_pref)   # never blank — falls back to "General"
            portfolio_1 = d.get("portfolio_1", "N/A")
            portfolio_2 = d.get("portfolio_2", "N/A")
            portfolio_3 = d.get("portfolio_3", "N/A")
            # Tier 1 gap-fill: any Portfolio slot still empty after regex gets one AI
            # attempt (a slot regex already found is never re-examined or second-guessed).
            portfolio_1, portfolio_2, portfolio_3 = infer_missing_portfolios(
                text, portfolio_1, portfolio_2, portfolio_3)
            email = d.get("email", "")
            location = d.get("location", "Not extracted")
            country = d.get("country", "")
            # Matches the SharePoint scoring path (sharepoint_scoring.py), which was the
            # only one of the two intake routes doing either of these until found live
            # 2026-09-01: a masked resume phone ('732-4***') came back as a fabricated-
            # looking '732-449***' instead of 'Not extracted', and a NANP-shaped number
            # (10 digit, or 11 starting with 1) was left as-written ('+1 480 930 2876')
            # instead of the uniform '(480) 930-2876'.
            phone = format_phone(d.get("phone", "Not extracted"))
            phone = sanitize_phone(phone)

            if GEO_FILTER_USA_ONLY:
                geo_decision, usa_reason = classify_location_usa(
                    location, country=country, resume_text=text, phone=phone,
                    education=d.get("education", "Not extracted"))
            else:
                geo_decision, usa_reason = GeoDecision.CONFIRMED_US, "geo filter off"
            row_status = (
                STATUS_SCORED
                if geo_decision == GeoDecision.CONFIRMED_US
                else STATUS_LOCATION_REVIEW
                if geo_decision == GeoDecision.UNKNOWN
                else _cfg.STATUS_REJECTED
            )

            # find an existing row
            match_idx, existing_name = None, ""
            for idx, row in df.iterrows():
                if _matches(name, email, row.get("Original Filename"), row.get("Email")):
                    match_idx, existing_name = idx, str(row.get("Full Name", ""))
                    break

            # Step A - snapshot all 22 columns' starting state before any write, so Step E
            # (below) can show exactly what changed. A brand-new candidate starts all-blank.
            if match_idx is not None:
                before_snapshot = {c: df.loc[match_idx, c] for c in COLUMNS}
            else:
                before_snapshot = {c: "" for c in COLUMNS}
            # One line naming which columns already held a value (a brand-new candidate
            # starts all-blank); the per-column filled/healed picture is the Diff line below.
            _prefilled = [c for c in COLUMNS if not _is_gap(before_snapshot[c])]
            logger.info("   Before  : " + (
                f"{len(_prefilled)} column(s) pre-filled: {', '.join(_prefilled)}"
                if _prefilled else "new candidate (all columns blank)"))

            if match_idx is not None:                       # ── UPDATE ──
                full = resolve_full_name(d.get("full_name"), existing_name, text)
                fields = {"Full Name": full, "Phone": phone, "Location": location,
                          "Country": country, "Current Skills": skills, "Looking For Role": role_pref,
                          "Suggested Role 1": r1, "Suggested Role 2": r2, "Suggested Role 3": r3,
                          "Category": category,
                          "Portfolio 1": portfolio_1, "Portfolio 2": portfolio_2,
                          "Portfolio 3": portfolio_3,
                          "Status": row_status}
                logger.info(f"   UPDATE  {name} -> {full} | {r1} | {r2} | {r3} "
                            f"[{res['source']}] | {time.perf_counter()-t0:.1f}s")
                logger.info(f"   location='{location}' | country='{country}' | "
                            f"geo={geo_decision.value} ({usa_reason})")
                _final = dict(before_snapshot); _final.update(fields)
                _log_column_diff(before_snapshot, _final, after_extract_snapshot)
                if geo_decision == GeoDecision.CONFIRMED_NON_US:
                    fields["Status"] = _cfg.STATUS_REJECTED
                    if not dry_run:
                        merged = df.loc[match_idx].to_dict()
                        merged.update(fields)
                        rejected_rows.append(merged)
                        df = df.drop(match_idx).reset_index(drop=True)
                    rejected += 1
                else:
                    if not dry_run:
                        for k, val in fields.items():
                            df.loc[match_idx, k] = val
                    updated += 1
            else:                                           # ── ADD (new) ──
                app_id = _new_app_id()
                full = resolve_full_name(d.get("full_name"), "", text)
                stored = f"{app_id}_{name}"
                logger.info(f"   ADD     {name} -> {app_id} | {full} | {r1} | {r2} | {r3} "
                            f"[{res['source']}] | {time.perf_counter()-t0:.1f}s")
                # 'geo=', not 'usa=': the boolean is_usa was replaced by the tri-state
                # GeoDecision (CONFIRMED_US / CONFIRMED_NON_US / UNKNOWN). The UPDATE branch
                # above was migrated, this ADD branch was missed, so every NEW candidate
                # reaching intake died on NameError: name 'is_usa' is not defined - the row
                # was scored correctly and then thrown away. Only --score-folder/--score-excel
                # run this path; production --score-sharepoint has its own, which is why it
                # survived unnoticed. Caught 2026-08-24 by a 3-candidate dry run.
                logger.info(f"   location='{location}' | country='{country}' | "
                            f"geo={geo_decision.value} ({usa_reason})")
                row_data = {"Application ID": app_id,
                            # ISO 'T' form: Excel Online keeps it as text (the space
                            # form gets coerced to a date serial; see excel_output).
                            "Received Date": now.strftime("%Y-%m-%dT%H:%M:%S"),
                            "Category": category,
                            "Full Name": full, "Email": email or "Not extracted",
                            "Phone": phone, "Location": location, "Country": country,
                            "Current Skills": skills, "Looking For Role": role_pref,
                            "Suggested Role 1": r1, "Suggested Role 2": r2, "Suggested Role 3": r3,
                            "Mail Subject": "", "Mail Body": "",
                            "Status": row_status, "Has Resume": "Yes",
                            "Original Filename": name, "Application Updates": 0,
                            "Portfolio 1": portfolio_1, "Portfolio 2": portfolio_2,
                            "Portfolio 3": portfolio_3}
                _log_column_diff(before_snapshot, row_data, after_extract_snapshot)
                if geo_decision == GeoDecision.CONFIRMED_NON_US:
                    row_data["Status"] = _cfg.STATUS_REJECTED
                    if not dry_run:
                        dest = _cfg.dated_dir(_cfg.INPUT_DIR, now)
                        dest.mkdir(parents=True, exist_ok=True)
                        (dest / stored).write_bytes(raw)
                        rejected_rows.append(row_data)
                    rejected += 1
                else:
                    if not dry_run:
                        dest = _cfg.dated_dir(_cfg.INPUT_DIR, now)
                        dest.mkdir(parents=True, exist_ok=True)
                        (dest / stored).write_bytes(raw)
                        df = pd.concat([df, pd.DataFrame([row_data], columns=COLUMNS)],
                                       ignore_index=True)
                    added += 1
        except Exception as e:
            errors += 1
            logger.error(f"   FAILED  {name}: {e}")

    if not dry_run:
        _save_local(df, target_xlsx)
        if rejected_rows:
            _save_rejected_local(rejected_rows, target_xlsx)
    logger.info(f"Intake done: {added} added, {updated} updated, {rejected} rejected (non-USA), "
                f"{skipped} skipped, {errors} error(s).")
    return {"added": added, "updated": updated, "rejected": rejected,
            "skipped": skipped, "errors": errors, "workbook": str(target_xlsx)}


def score_existing_excel(excel_file, dry_run: bool = False) -> dict:
    """Run the scoring brain over the rows of an existing workbook, in place.

    Re-scores 'Suggested Role 1/2/3' + 'Category' for every row from its 'Current Skills'
    and 'Looking For Role'. Returns a summary dict.
    """
    from hiring_agent.scoring import rematch_sheet
    target = Path(excel_file)
    if not target.exists():
        logger.error(f"Workbook not found: {target}")
        return {"error": f"Workbook not found: {target}"}
    roles = get_active_roles()
    logger.info(f"Score-Excel | workbook='{target.name}' | {'DRY-RUN' if dry_run else 'LIVE'}")
    if dry_run:
        df = pd.read_excel(target)
        for _, row in df.iterrows():
            res = suggested_roles(str(row.get("Current Skills", "")),
                                  str(row.get("Looking For Role", "")), roles=roles)
            logger.info(f"   [DRY] {row.get('Full Name', '?')}: "
                        f"{res['role_1']} | {res['role_2']} [{res['source']}]")
        return {"rows": len(df), "workbook": str(target), "dry_run": True}
    rematch_sheet(excel_file=target, roles=roles)
    return {"workbook": str(target)}
