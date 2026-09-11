"""Excel Report Exporter for P2 Hiring Agent.

Generates pristine `Candidate_List_Results.xlsx` from CandidateStore (SQLite/Excel)
and uploads it atomically to SharePoint via single PUT content.
"""
from __future__ import annotations

import datetime
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from hiring_agent.config import COLUMNS, REJECTED_COLUMNS, logger
from hiring_agent.store import CandidateStore, CandidateRow


class ExcelReportExporter:
    """Generates formatted Excel reports from CandidateStore."""

    def __init__(self, store: CandidateStore):
        self.store = store

    def generate_workbook(self, output_path: str | Path) -> Path:
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        from hiring_agent.config import STATUS_SCORED
        from hiring_agent.sharepoint_scoring import (
            _client_export_sort_key,
            prepare_client_export_rows,
            write_client_export,
            audit_client_export_integrity,
        )

        main_candidates = self.store.all_rows("main")
        scored_candidates = []
        for row in main_candidates:
            vals = row.values if isinstance(row, CandidateRow) else (row.get("values", {}) if isinstance(row, dict) else {})
            if not str(vals.get("Application ID", "") or "").strip():
                continue
            if str(vals.get("Status", "") or "").strip() != STATUS_SCORED:
                continue
            country = str(vals.get("Country", "") or "").strip().lower()
            if country and country not in ("united states", "usa", "us"):
                continue
            phone = str(vals.get("Phone", "") or "").strip()
            if phone.startswith("+") and not (phone.startswith("+1") or phone.startswith("+ 1")):
                continue
            scored_candidates.append(vals)

        scored_candidates.sort(key=_client_export_sort_key, reverse=True)

        seen_applicants = set()
        deduped = []
        for vals in scored_candidates:
            email = str(vals.get("Email", "") or "").strip().lower()
            full_name = str(vals.get("Full Name", "") or "").strip().lower()
            key = email if (email and "@" in email) else (full_name if full_name else vals.get("Application ID"))
            if key in seen_applicants:
                continue
            seen_applicants.add(key)
            deduped.append(vals)

        prepared = prepare_client_export_rows(deduped)
        write_client_export(prepared, output_path)
        return output_path

    def _populate_sheet(self, ws, sheet_name: str, columns: list[str]):
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        cell_font = Font(name="Calibri", size=11)
        align_left = Alignment(horizontal="left", vertical="center", wrap_text=False)
        border_thin = Side(border_style="thin", color="D9D9D9")
        cell_border = Border(top=border_thin, bottom=border_thin, left=border_thin, right=border_thin)

        # Write Header
        ws.append(columns)
        for col_idx in range(1, len(columns) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = align_left

        # Write Rows
        rows = self.store.all_rows(sheet=sheet_name)
        for row_idx, r in enumerate(rows, start=2):
            vals = r.values if isinstance(r, CandidateRow) else (r.get("values", {}) if isinstance(r, dict) else {})
            row_data = []
            for col in columns:
                if col == "Resume Link":
                    url = str(vals.get("Resume URL", "") or "").strip()
                    fn = str(vals.get("Original Filename", "") or "").strip() or "Resume"
                    if url and url.startswith("http"):
                        formula = f'=HYPERLINK("{url}", "{fn}")'
                    else:
                        formula = vals.get(col, "")
                    row_data.append(formula)
                else:
                    row_data.append(vals.get(col, ""))
            ws.append(row_data)

            # Apply cell styles
            for col_idx, col in enumerate(columns, start=1):
                c = ws.cell(row=row_idx, column=col_idx)
                c.font = cell_font
                c.alignment = align_left
                c.border = cell_border
                if col in ("Phone", "Education Start Date", "Education End Date"):
                    c.number_format = "@"

        # Auto-fit column widths
        for col in ws.columns:
            vals_col = [str(cell.value or "") for cell in col]
            max_len = max(len(v) for v in vals_col) if vals_col else 10
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    def export_and_upload(self, client=None, local_dir: str | Path = "P2_Final_Results") -> tuple[Path, bool]:
        """Generate workbook, save locally, and upload to SharePoint atomically."""
        local_dir = Path(local_dir)
        local_dir.mkdir(parents=True, exist_ok=True)
        ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        latest_file = local_dir / "Candidate_List_Results.xlsx"
        backup_file = local_dir / f"Candidate_List_Results_{ts}.xlsx"

        self.generate_workbook(latest_file)
        import shutil
        shutil.copyfile(latest_file, backup_file)
        logger.info(f"   Exported results to {latest_file} (and backup {backup_file.name})")

        uploaded = False
        if client and hasattr(client, "upload_file") and hasattr(client, "_wb_path"):
            try:
                import os
                from .publisher import safe_target
                target = os.getenv("HIRING_CLIENT_REPORT_PATH", "/Candidate_List_Results.xlsx")
                safe_target(target, client._wb_path)
                folder, _, filename = target.strip('/').rpartition('/')
                data = latest_file.read_bytes()
                client.upload_file(folder, filename, data)
                uploaded = True
                logger.info(f"   Uploaded {filename} atomically to SharePoint folder '{folder}'")
            except Exception as e:
                logger.warning(f"   Could not upload {latest_file.name} to SharePoint: {e}")

        return latest_file, uploaded
